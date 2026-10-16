"""Corrected I/O allocation for the Schneider Electric requote (Modicon M580 + Edge I/O NTS).

Allocates every hardwired signal of the Villa Somalia Device and I/O List Rev 1 to a
Schneider module and channel at the location where the signal is wired, and lists the
changes against Schneider's offer (Offer Proposal Villa Somalia, version 1, 16 Oct 2026).

Usage: python3 tools/schneider_io_allocation.py <io_list_rev1.xlsx> <out.xlsx>
"""
import sys
from collections import OrderedDict
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]

# ------------------------------------------------------------------ source signals
io = openpyxl.load_workbook(SRC)['IO List']
hdr = [c.value for c in io[1]]
col = {h: i for i, h in enumerate(hdr)}
sigs = [r for r in io.iter_rows(min_row=2, values_only=True) if r[0]]
assert len(sigs) == 244, len(sigs)

# ------------------------------------------------------------------ Schneider modules per location
DDI32, DDO32, AMI4 = 'BMXDDI3202K', 'BMXDDO3202K', 'BMXAMI0410'
NDI16, NDO16, NAI4 = 'NTSDDI1602XK', 'NTSDDO1602K', 'Edge I/O NTS analogue input, 4 x 4-20 mA (Schneider to select)'
MODDESC = {
    DDI32: 'X80 digital input, 32 channels, 24 V DC', DDO32: 'X80 digital output, 32 channels, 24 V DC',
    AMI4: 'X80 analogue input, 4 isolated channels', NDI16: 'Edge I/O NTS digital input, 16 channels, 24 V DC',
    NDO16: 'Edge I/O NTS digital output, 16 channels, 24 V DC, 0.5 A', NAI4: 'Edge I/O NTS analogue input, 4 channels',
}
LOC = OrderedDict([
    ('SWR', ('11 kV switchroom, M580 Ethernet remote drop (BMECRA31210 on BMEXBP1200)', '1, 2')),
    ('PRS', ('President office station, Edge I/O island', '3')),
    ('PMS', ('Prime Minister office station, Edge I/O island', '4')),
    ('YRD', ('BESS and inverter transformer yard, Edge I/O island', '5')),
    ('GEN', ('Generator area, Edge I/O island', '6')),
])
MODULES = OrderedDict()  # module id -> (location, reference, channels)
for i in range(1, 5): MODULES['SWR-DI%d' % i] = ('SWR', DDI32, 32)
MODULES['SWR-DO1'] = ('SWR', DDO32, 32)
MODULES['SWR-AI1'] = ('SWR', AMI4, 4)
for st in ('PRS', 'PMS'):
    for i in range(1, 4): MODULES['%s-DI%d' % (st, i)] = (st, NDI16, 16)
    MODULES['%s-DO1' % st] = (st, NDO16, 16)
for i in range(1, 3): MODULES['YRD-DI%d' % i] = ('YRD', NDI16, 16)
for i in range(1, 3): MODULES['GEN-DI%d' % i] = ('GEN', NDI16, 16)
MODULES['GEN-AI1'] = ('GEN', NAI4, 4)


def split(items, n):
    """Spread items evenly over n modules so every module keeps some spare channels."""
    k, m = divmod(len(items), n)
    out, p = [], 0
    for i in range(n):
        q = k + (1 if i < m else 0)
        out.append(items[p:p + q]); p += q
    return out


def of(row, typ, rev_mod=None):
    return [s for s in sigs if s[col['I/O row']] == row and s[col['Type']] == typ and (rev_mod is None or s[col['Module type']] == rev_mod)]


alloc = []  # (signal, module id, channel, interface, relay)
relay_no = [0]


def place(mod, items, start=1, cmd=False):
    for ch, s in enumerate(items, start):
        if s[col['Type']] == 'DI':
            iface = 'Dry contact, wetted 24 V DC from the panel supply'
            relay = ''
        elif s[col['Type']] == 'AI':
            iface = '4-20 mA, ' + str(s[col['Normal state / range']])
            relay = ''
        else:
            relay_no[0] += 1
            relay = 'IR%02d' % relay_no[0]
            iface = ('Interposing relay %s, contact in the %s control circuit' % (relay, 'breaker or switch' if cmd else 'beacon or watchdog'))
        alloc.append((s, mod, ch, iface, relay))


# switchroom: row 1 inputs over DI1-DI3, row 2 inputs on DI4; commands ch 1-14, beacon and watchdog ch 17-18
for mod, part in zip(('SWR-DI1', 'SWR-DI2', 'SWR-DI3'), split(of(1, 'DI'), 3)): place(mod, part)
place('SWR-DI4', of(2, 'DI'))
place('SWR-DO1', of(1, 'DO', 'DO-8221'), cmd=True)
place('SWR-DO1', of(1, 'DO', 'DO-8212'), start=17)
place('SWR-AI1', of(2, 'AI'))
for st, row in (('PRS', 3), ('PMS', 4)):
    for mod, part in zip(('%s-DI1' % st, '%s-DI2' % st, '%s-DI3' % st), split(of(row, 'DI'), 3)): place(mod, part)
    place('%s-DO1' % st, of(row, 'DO'), cmd=True)
for mod, part in zip(('YRD-DI1', 'YRD-DI2'), split(of(5, 'DI'), 2)): place(mod, part)
for mod, part in zip(('GEN-DI1', 'GEN-DI2'), split(of(6, 'DI'), 2)): place(mod, part)
place('GEN-AI1', of(6, 'AI'))
assert len(alloc) == 244 and len({(a[1], a[2]) for a in alloc}) == 244
N_RELAYS = relay_no[0]

# ------------------------------------------------------------------ styles
NAVY, PALE, LINE, YEL = '1F3A5F', 'E8EEF5', 'BFC8D2', 'FFF4C2'
thin = Side(style='thin', color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HFONT, HFILL = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor=NAVY)
WRAP = Alignment(wrap_text=True, vertical='top')


def table(ws, r0, head, rows, widths=None):
    for c, h in enumerate(head, 1):
        x = ws.cell(row=r0, column=c, value=h)
        x.font, x.fill, x.border, x.alignment = HFONT, HFILL, BOX, Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            x = ws.cell(row=i, column=c, value=v)
            x.border, x.alignment = BOX, WRAP
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    return r0 + len(rows)


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = Font(italic=True, color='5B5B57')


wb = openpyxl.Workbook()

# ------------------------------------------------------------------ Read Me
rm = wb.active
rm.title = 'Read Me'
title(rm, 'Villa Somalia SCADA: corrected I/O allocation for the Schneider Electric requote')
info = [
    ('Purpose', 'Basis for Schneider Electric to requote the PLC and remote I/O. Allocates each hardwired signal of the Device and I/O List Rev 1 to a Schneider module and channel at the location where it is wired, and states the changes against the offer.'),
    ('Offer reviewed', 'Offer Proposal Villa Somalia, version 1, 16 Oct 2026 (Schneider Electric), and its Bill of Material workbook.'),
    ('Signal basis', 'Villa_Somalia_SCADA_Device_and_IO_List_Rev1.xlsx, sheet IO List: 244 hardwired signals (208 DI, 32 DO, 4 AI). The 463 communication points (sheet Comms Point List) are read over the network and need no I/O channels.'),
    ('Why the offer does not fit', 'The offer has enough channels in total (416 DI, 192 DO) but they are in the wrong places: four PV inverter islands where the yard needs 26 inputs, no I/O in the generator area, no analogue inputs, both office islands labelled PM office, six times the outputs required, and no interposing relays for the breaker commands.'),
    ('Locations', 'Switchroom (I/O rows 1 and 2 of the design), President station (row 3), Prime Minister station (row 4), BESS and inverter yard (row 5), generator area (row 6).'),
    ('Switchroom I/O', 'All switchroom I/O sits in the Ethernet remote drop, none in the CPU rack, so the CPU rack can become a hot standby pair without rewiring (sheet Other Items, item 1).'),
    ('Spare capacity', 'About 20 % or more spare digital inputs per location, as in the design: 27 % in the switchroom and office stations, 44 % in the generator area, 19 % in the BESS yard (add a third input module there if more yard signals appear). Outputs and analogue inputs are single modules with 50 % spare.'),
    ('Inputs', 'Switchgear, RMU, transformer and package contacts are assumed dry contacts, wetted at 24 V DC from a field supply in each panel. Contacts that carry 110 V DC need an interposing relay; confirm with the terminal schedules.'),
    ('Outputs', 'Breaker and switch commands drive interposing relays (%d in total: %d commands, beacon and watchdog). The relay contact switches the 110 V DC control circuit; the 24 V DC transistor outputs never switch a coil directly.' % (N_RELAYS, N_RELAYS - 2)),
    ('Sheets', 'Hardware by Location: modules per location with used and spare channels, calculated. IO Allocation: every signal with its Schneider module and channel. Changes to Offer: line-by-line quantities. Other Items: what else must be quoted or clarified.'),
    ('Status', 'Rev 0, 16 Oct 2026, for Schneider Electric to requote. Module and channel numbers are proposals; Schneider may re-number within a location if the totals and spare stay the same.'),
]
r = table(rm, 4, ['Item', 'Description'], info, [24, 120])
for i in range(5, r + 1):
    rm.row_dimensions[i].height = 48

# ------------------------------------------------------------------ IO Allocation
al = wb.create_sheet('IO Allocation')
head = ['Tag', 'Description', 'Equipment', 'Type', 'Normal state / range', 'Alarm level', 'Operated by',
        'Design row', 'Design module', 'Design channel', 'Schneider location', 'Schneider module', 'Module reference', 'Channel', 'Interface', 'Interposing relay']
rows = []
for s, mod, ch, iface, relay in alloc:
    loc, ref, _ = MODULES[mod]
    rows.append([s[col['Tag']], s[col['Description']], s[col['Equipment']], s[col['Type']], s[col['Normal state / range']], s[col['Alarm level']],
                 s[col['Operated by']], s[col['I/O row']], s[col['Module']], s[col['Channel']], loc, mod, ref if ref != NAI4 else 'NTS analogue input (to select)', ch, iface, relay])
last = table(al, 1, head, rows, [22, 40, 34, 6, 16, 8, 10, 8, 10, 8, 10, 11, 18, 8, 46, 10])
al.freeze_panes = 'B2'
al.auto_filter.ref = 'A1:P%d' % last
for i in range(2, last + 1):
    al.row_dimensions[i].height = 15
    for c in range(1, 17):
        al.cell(row=i, column=c).alignment = Alignment(vertical='top')
RNG = "'IO Allocation'!$L$2:$L$%d" % last
TRNG = "'IO Allocation'!$K$2:$K$%d" % last
TYP = "'IO Allocation'!$D$2:$D$%d" % last

# ------------------------------------------------------------------ Hardware by Location
hw = wb.create_sheet('Hardware by Location', 1)
title(hw, 'I/O hardware by location', 'Used channels are counted from sheet IO Allocation.')
head = ['Location', 'Design rows', 'Module', 'Module reference', 'Description', 'Channels', 'Used', 'Spare', 'Spare %']
r0 = 4
for c, h in enumerate(head, 1):
    x = hw.cell(row=r0, column=c, value=h)
    x.font, x.fill, x.border, x.alignment = HFONT, HFILL, BOX, Alignment(wrap_text=True)
r = r0 + 1
first_mod_row = r
for mid, (loc, ref, ch) in MODULES.items():
    vals = [LOC[loc][0], LOC[loc][1], mid, ref if ref != NAI4 else 'To select', MODDESC[ref], ch,
            '=COUNTIF(%s,C%d)' % (RNG, r), '=F%d-G%d' % (r, r), '=IF(F%d=0,0,H%d/F%d)' % (r, r, r)]
    for c, v in enumerate(vals, 1):
        x = hw.cell(row=r, column=c, value=v)
        x.border, x.alignment = BOX, WRAP
    hw.cell(row=r, column=9).number_format = '0%'
    r += 1
last_mod_row = r - 1
r += 1
hw.cell(row=r, column=1, value='Totals by location and type').font = Font(bold=True, color=NAVY)
r += 1
head2 = ['Location', 'Type', 'Signals needed', 'Channels installed', 'Spare', 'Spare %']
for c, h in enumerate(head2, 1):
    x = hw.cell(row=r, column=c, value=h)
    x.font, x.fill, x.border = HFONT, HFILL, BOX
r += 1
tot_rows = []
for loc in LOC:
    for typ in ('DI', 'DO', 'AI'):
        mods = [m for m, v in MODULES.items() if v[0] == loc and m.split('-')[1].startswith(typ)]
        if not mods:
            continue
        chans = sum(MODULES[m][2] for m in mods)
        vals = [LOC[loc][0].split(',')[0], typ, '=COUNTIFS(%s,"%s",%s,"%s")' % (TRNG, loc, TYP, typ), chans, '=D%d-C%d' % (r, r), '=IF(D%d=0,0,E%d/D%d)' % (r, r, r)]
        for c, v in enumerate(vals, 1):
            x = hw.cell(row=r, column=c, value=v)
            x.border = BOX
        hw.cell(row=r, column=6).number_format = '0%'
        tot_rows.append(r)
        r += 1
for typ in ('DI', 'DO', 'AI'):
    rs = [x for x in tot_rows if hw.cell(row=x, column=2).value == typ]
    vals = ['All locations', typ, '=' + '+'.join('C%d' % x for x in rs), '=' + '+'.join('D%d' % x for x in rs), '=D%d-C%d' % (r, r), '=IF(D%d=0,0,E%d/D%d)' % (r, r, r)]
    for c, v in enumerate(vals, 1):
        x = hw.cell(row=r, column=c, value=v)
        x.border, x.font = BOX, Font(bold=True)
    hw.cell(row=r, column=6).number_format = '0%'
    r += 1
r += 1
notes = [
    'Switchroom: one M580 Ethernet remote drop (BMEXBP1200, BMXCPS3020, BMECRA31210) holds all switchroom I/O, the analogue module and the BMXNOM0200 serial module; the CPU rack holds no I/O.',
    'Each Edge I/O island: NTSNEC1200K network interface, NTSPFB1002HK power supply, the modules listed, terminal blocks per module, in an enclosure suited to its location (outdoor IP55 at the yard and generator area).',
    'Office station islands also carry an NTSFMB0120K serial master for the ST260E RMU relay (Modbus RTU). Add one at the generator area if the generator controllers are RS485 only.',
    'President and Prime Minister islands are identical; the offer labels both "PM office LVDB".',
]
for t in notes:
    hw.cell(row=r, column=1, value=t).alignment = Alignment(wrap_text=True, vertical='top')
    hw.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    hw.row_dimensions[r].height = 30
    r += 1
for c, w in enumerate([44, 9, 11, 16, 44, 9, 8, 8, 8], 1):
    hw.column_dimensions[get_column_letter(c)].width = w
hw.freeze_panes = 'A5'

# ------------------------------------------------------------------ Changes to Offer
ch = wb.create_sheet('Changes to Offer')
title(ch, 'Changes against the Schneider offer', 'Quantities include one spare per module type, as in the offer. "Not priced" = in the offer PDF and drawing but not in the priced bill of material.')
CH = [
    # reference, description, offer qty, proposed qty, change, reason
    ('HMIP67GD0F88BE7N00', 'Industrial panel PC (no description in the offer)', 2, 2, 'Keep; clarify', 'Confirm model, and whether the SCADA runs on these PCs or on servers (Other Items, item 4).'),
    ('HMIEMSE3PRT4KA', 'Machine SCADA Expert runtime, 4K tags', 2, 2, 'Keep', '707 tags (244 hardwired + 463 communication) fit in 4K.'),
    ('HMIEMSEBT4KA', 'Machine SCADA Expert build-time, 4K tags', 1, 1, 'Keep', ''),
    ('BMEXBP0800', 'X80 8-slot Ethernet backplane (CPU rack)', 2, 2, 'Keep', 'CPU rack, no I/O. See Other Items, item 1 for hot standby.'),
    ('BMXCPS3020', 'X80 power supply 24-48 V DC', 3, 3, 'Keep', 'CPU rack, switchroom drop, spare.'),
    ('BMEP582040', 'M580 CPU', 2, 2, 'Keep or replace', 'Simplex as offered; the design requires a hot standby pair (Other Items, item 1).'),
    ('BMENOC0311', 'M580 Ethernet communication module', 2, 2, 'Keep', ''),
    ('BMENOC0321', 'M580 Ethernet module, IP forwarding', 1, 1, 'Keep', ''),
    ('BMEXBP1200', 'X80 12-slot Ethernet backplane (switchroom drop)', 1, 1, 'Keep', 'Holds CRA, 4 DI, 1 DO, 1 AI and the serial module.'),
    ('BMECRA31210', 'X80 Ethernet drop adapter', 1, 1, 'Keep', ''),
    (DDI32, MODDESC[DDI32], 5, 5, 'Same quantity, moved', 'All 4 in the switchroom drop (94 inputs on 128 channels) + 1 spare; none in the CPU rack.'),
    (DDO32, MODDESC[DDO32], 4, 2, 'Reduce by 2', 'Switchroom has 16 outputs: 1 module + 1 spare.'),
    ('BMXFCW303', 'Pre-wired cable for 32-channel modules', 7, 7, 'Keep', 'One per 32-channel module including spares.'),
    (AMI4, MODDESC[AMI4] + ', with its terminal block', 0, 2, 'Add 2', 'DC bus voltages 110 V and 24 V (4-20 mA); 1 + 1 spare.'),
    ('BMXNOM0200', 'X80 serial module, Modbus RTU', 1, 1, 'Keep; confirm use', 'State which switchroom devices it serves.'),
    ('VW3A8306R10', 'RS485 cable', 1, 1, 'Keep', ''),
    ('MCSESM083F23F1', 'Managed switch, 8 copper ports', 2, 2, 'Clarify', 'No fibre ports. Remote locations need fibre (Other Items, item 2).'),
    ('NTSNEC1200K', 'Edge I/O network interface', 7, 5, 'Reduce by 2', '4 islands (President, PM, yard, generators) + 1 spare.'),
    ('NTSPFB1002HK', 'Edge I/O power supply module', 7, 5, 'Reduce by 2', 'As above.'),
    ('NTSXTB02030H', 'Terminal block for the power supply module', 12, 8, 'Reduce by 4', '2 per island.'),
    (NDI16, MODDESC[NDI16], 19, 11, 'Reduce by 8', 'President 3, PM 3, yard 2, generators 2, + 1 spare.'),
    ('NTSXTB18001XH', 'Terminal block for NTSDDI1602XK', 18, 10, 'Reduce by 8', 'One per input module in use.'),
    (NDO16, MODDESC[NDO16], 7, 3, 'Reduce by 4', 'President 1, PM 1 (8 commands each), + 1 spare.'),
    ('NTSXTB18001H', 'Terminal block for NTSDDO1602K', 6, 2, 'Reduce by 4', 'One per output module in use.'),
    ('NTSFMB0120K', 'Edge I/O serial master, Modbus RTU', 6, 2, 'Reduce by 4', 'ST260E relay at each office station. Add 1 at the generators only if their controllers are RS485.'),
    ('NTSXTB12001H', 'Terminal block for NTSFMB0120K', 6, 2, 'Reduce by 4', ''),
    ('To select', 'Edge I/O NTS analogue input, 4 x 4-20 mA, with terminal block', 0, 2, 'Add 2', 'Generator fuel levels DG1 and DG2; 1 + 1 spare.'),
    ('To select', 'Interposing relay, 24 V DC coil, contact rated for the 110 V DC control circuit, with base', 0, N_RELAYS + 2, 'Add %d' % (N_RELAYS + 2),
     '%d command relays (AH1-AH7 open/close, RMU and LV incomer commands at both stations) + beacon and watchdog + 2 spare.' % (N_RELAYS - 2)),
    ('To select', '24 V DC field supply for contact wetting, I/O and the X80/NTS power supplies', 0, 5, 'Add 5', 'One per location; supplied from the 110 V DC station battery where available.'),
    ('PVSCL60E', 'Schneider three-phase string inverter (not priced)', 4, 0, 'Remove', 'The plant has 4 Sungrow SG350HX inverters, read through the PLC box over Modbus TCP.'),
    ('PAS600 + CL110 + TH110', 'Panel servers with humidity and temperature sensors', 7, 0, 'Move to option', '7 panel servers, 7 humidity and 9 temperature sensors are not in the I/O list. Quote separately as an option for the client to decide.'),
    ('REL50458', 'Easergy P5M30 protection relay (not priced)', 8, 0, 'Clarify', 'The 11 kV relays may come with the switchgear. If Schneider supplies them, the P5M30 appears to be the motor variant; feeders normally use P5F30.'),
    ('METSEPM5320', 'PM5320 power meter (not priced)', 5, 0, 'Clarify', 'The design has 3 x PM5560 (President LVDB, PM LVDB, auxiliary board). State where 5 meters would go.'),
    ('Lot', 'System design and commissioning', 1, 1, 'Requote', 'To the corrected scope.'),
]
r = table(ch, 4, ['Reference', 'Description', 'Offer qty', 'Proposed qty', 'Change', 'Reason'], CH, [22, 52, 9, 11, 18, 70])
for i in range(5, r + 1):
    ch.row_dimensions[i].height = 30
    if str(ch.cell(row=i, column=5).value).startswith(('Add', 'Remove', 'Clarify', 'Move')):
        ch.cell(row=i, column=5).fill = PatternFill('solid', fgColor=YEL)
ch.freeze_panes = 'A5'

# ------------------------------------------------------------------ Other Items
ot = wb.create_sheet('Other Items')
title(ot, 'Other items to quote or clarify', 'Requirements from the SCADA Design Report Rev 1 that the offer does not cover.')
OT = [
    (1, 'Controller redundancy', 'The design requires a hot standby controller pair (design report section 4). Quote M580 Hot Standby (two BMEH58 processors on hot standby racks, as described in the offer) with all I/O in the switchroom drop and the Edge I/O islands. State the switchover time. Keep the simplex CPU as a price option.'),
    (2, 'Fibre network', 'The office stations, BESS yard and generator area are beyond the 100 m copper limit. Quote managed switches with SFP fibre ports in a ring, or state how the PLC and Edge I/O connect to the fibre ring of the design.'),
    (3, 'Cybersecurity', 'Firewall with VPN for remote support (off by default), and time synchronisation (GPS/NTP). State the M580 and SCADA security settings used.'),
    (4, 'SCADA', 'State whether the SCADA is redundant, where the historian runs and its retention, and confirm Operator and Admin roles, three alarm levels, the event log with user names, and report export.'),
    (5, 'Communication interfaces', 'Confirm each interface and its points per sheet Comms Point List of the I/O list: 11 kV relays (IEC 61850 or Modbus TCP), EMS (Modbus TCP), PLC box (Modbus TCP), ST260E (Modbus RTU), power meters, weather logger, UPS, network devices (SNMP).'),
    (6, 'Control authority', 'The embedded EMS dispatches PV, battery and generators. The PLC monitors, checks interlocks on remote commands and passes Admin settings to the EMS; it does not run plant dispatch sequences. Correct the offer text accordingly.'),
    (7, 'Emergency stop', 'Hardwired only; the PLC and SCADA only display it. No software emergency stop.'),
    (8, 'Enclosures and wiring', 'State who supplies the Edge I/O enclosures (outdoor IP55 at the yard and generator area), marshalling, interposing relay panels and field wiring.'),
    (9, 'Commercial terms', 'State currency, validity, delivery time, payment terms, warranty and exclusions. The General Conditions section of the offer is empty.'),
    (10, 'Offer text', 'Remove text from another project (treated-water process, IICATS, LIMS, Quantum RIO), complete the revision history and table of contents, and correct the second office island to the President station.'),
]
r = table(ot, 4, ['No.', 'Item', 'Requirement'], OT, [6, 26, 120])
for i in range(5, r + 1):
    ot.row_dimensions[i].height = 45

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('signals', len(alloc), 'modules', len(MODULES), 'relays', N_RELAYS, '->', OUT)
