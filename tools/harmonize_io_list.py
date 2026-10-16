"""Rev 1 of the Villa Somalia SCADA Device and I/O List.

Adds the signals found missing in the I/O review (AH8 breaker, LV main breakers,
inverter AC MCCBs, auxiliary transformer fuse, RMU UPS, relay trip-circuit /
SEF / reverse-power points, battery combiner panels), makes the CP-8050 a hot
standby pair (controller A active, B standby; I/O row 1 moved off the controller
backplane so both controllers reach it), keeps every formula range
in step, and adds a 'Comms Point List' sheet that expands each communication
point line into one tagged signal per device.

Usage: python3 tools/harmonize_io_list.py <rev0.xlsx> <out.xlsx>
"""
import sys, re
from copy import copy
import openpyxl

SRC, OUT = sys.argv[1], sys.argv[2]
REV = 'Added rev 1'
wb = openpyxl.load_workbook(SRC)


def read_rows(ws, first, last):
    rows = []
    for r in range(first, last + 1):
        cells = [ws.cell(row=r, column=c) for c in range(1, ws.max_column + 1)]
        rows.append(([c.value for c in cells], [copy(c._style) for c in cells]))
    return rows


def write_rows(ws, first, rows, clear_to):
    for i, (vals, styles) in enumerate(rows):
        r = first + i
        for c, (v, st) in enumerate(zip(vals, styles), 1):
            cell = ws.cell(row=r, column=c)
            cell.value = v
            cell._style = copy(st)
    for r in range(first + len(rows), clear_to + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None


def insert_after(rows, key_col, key, new_vals):
    """Insert new rows after the row whose column key_col equals key, using its style."""
    idx = next(i for i, (v, _) in enumerate(rows) if v[key_col] == key)
    style = rows[idx][1]
    width = len(rows[idx][0])
    for k, nv in enumerate(new_vals):
        nv = list(nv) + [None] * (width - len(nv))
        rows.insert(idx + 1 + k, (nv, [copy(s) for s in style]))


def replace_in_formulas(ws, old, new):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and old in c.value:
                c.value = c.value.replace(old, new)


# ---------------------------------------------------------------- IO List
io = wb['IO List']
last = max(r for r in range(2, io.max_row + 1) if io.cell(row=r, column=1).value)
rows = read_rows(io, 2, last)
loc = {v[0]: v[3] for v, _ in rows}


def di(tag, desc, equip, location, row, module, ch, normal, alarm, note=REV):
    return [tag, desc, equip, location, 'DI', row, module, 'DI-8112', ch, normal, alarm, '-', note]


AH8 = '11 kV panel AH8, neutral earthing resistor'
insert_after(rows, 0, 'AH8-AUX-OK', [
    di('AH8-CB-OPEN', 'Circuit breaker open', AH8, loc['AH8-AUX-OK'], 1, 'R1-DI05', 3, 'Status', 'Status',
       REV + '. The 11 kV system is earthed only through AH8'),
    di('AH8-CB-CLSD', 'Circuit breaker closed', AH8, loc['AH8-AUX-OK'], 1, 'R1-DI05', 4, 'Status', 'Status',
       REV + '. Level 1 alarm if open while the busbar is live'),
    di('AH8-CB-TRIP', 'Circuit breaker tripped by protection', AH8, loc['AH8-AUX-OK'], 1, 'R1-DI05', 5, 'Normally open', '1'),
    di('AH8-ES-CLSD', 'Earth switch closed', AH8, loc['AH8-AUX-OK'], 1, 'R1-DI05', 6, 'Status', 'Status'),
    di('AH8-GAS-LOW', 'Gas pressure low', AH8, loc['AH8-AUX-OK'], 1, 'R1-DI05', 7, 'Normally open', '2'),
])
insert_after(rows, 0, 'T-AUX-OL-LOW', [
    di('T-AUX-FUSE', 'HV fuse blown (XRNT-3A striker)', 'Auxiliary transformer 100 kVA', loc['T-AUX-OL-LOW'], 2, 'R2-DI02', 7,
       'Normally open', '1', REV + '. The fuse is the only HV protection of this transformer'),
])
for st, name, row in (('PRS', 'President', 3), ('PMS', 'Prime Minister', 4)):
    eq = name + ' station RMU'
    insert_after(rows, 0, st + '-PNL-SUP-OK', [
        di(st + '-UPS-ON-BATT', 'RMU UPS 2 kVA on battery', eq, loc[st + '-PNL-SUP-OK'], row, 'R%d-DI03' % row, 2, 'Normally open', '2',
           REV + '. UPS shown on the transformer stations drawing'),
        di(st + '-UPS-FLT', 'RMU UPS 2 kVA fault', eq, loc[st + '-PNL-SUP-OK'], row, 'R%d-DI03' % row, 3, 'Normally open', '2'),
    ])
yard = loc['T-INV-OL-LOW']
PVLV, BLV = 'PV LV switchboard 0.8 kV (AC combiner)', 'BESS LV breaker 0.69 kV'
insert_after(rows, 0, 'T-INV-OL-LOW', [
    di('PV-ACB-CLSD', 'PV main LV breaker (ACB 1250 A) closed', PVLV, yard, 5, 'R5-DI02', 3, 'Status', 'Status'),
    di('PV-ACB-TRIP', 'PV main LV breaker (ACB 1250 A) tripped', PVLV, yard, 5, 'R5-DI02', 4, 'Normally open', '1'),
    di('NB01-MCCB-TRIP', 'Inverter NB01 AC breaker (MCCB 400 A) tripped', PVLV, yard, 5, 'R5-DI02', 5, 'Normally open', '2'),
    di('NB02-MCCB-TRIP', 'Inverter NB02 AC breaker (MCCB 400 A) tripped', PVLV, yard, 5, 'R5-DI02', 6, 'Normally open', '2'),
    di('NB03-MCCB-TRIP', 'Inverter NB03 AC breaker (MCCB 400 A) tripped', PVLV, yard, 5, 'R5-DI02', 7, 'Normally open', '2'),
    di('NB04-MCCB-TRIP', 'Inverter NB04 AC breaker (MCCB 400 A) tripped', PVLV, yard, 5, 'R5-DI02', 8, 'Normally open', '2'),
    di('BESS-ACB-CLSD', 'Battery main LV breaker (ACB 1000 A) closed', BLV, yard, 5, 'R5-DI02', 9, 'Status', 'Status'),
    di('BESS-ACB-TRIP', 'Battery main LV breaker (ACB 1000 A) tripped', BLV, yard, 5, 'R5-DI02', 10, 'Normally open', '1'),
])
insert_after(rows, 0, 'T-DG-OL-LOW', [
    di('GEN-ACB-CLSD', 'DG step-up transformer LV breaker closed', 'Generator step-up transformer', loc['T-DG-OL-LOW'], 6, 'R6-DI02', 1, 'Status', 'Status'),
    di('GEN-ACB-TRIP', 'DG step-up transformer LV breaker tripped', 'Generator step-up transformer', loc['T-DG-OL-LOW'], 6, 'R6-DI02', 2, 'Normally open', '1'),
])

tags = [v[0] for v, _ in rows]
assert len(tags) == len(set(tags)), 'duplicate tag'
chans = [(v[6], v[8]) for v, _ in rows]
assert len(chans) == len(set(chans)), 'channel used twice'
new_last = 1 + len(rows)
write_rows(io, 2, rows, max(last, new_last))
io.auto_filter.ref = 'A1:M%d' % new_last
for name in ('IO Summary', 'IO Rows'):
    replace_in_formulas(wb[name], '$%d' % last, '$%d' % new_last)

# ---------------------------------------------------------------- Comms Points
cp = wb['Comms Points']
cp_last = max(r for r in range(2, cp.max_row + 1) if cp.cell(row=r, column=1).value and cp.cell(row=r, column=3).value)
foot = read_rows(cp, cp_last + 1, cp.max_row)
crows = read_rows(cp, 2, cp_last)
relay_src = next(v[1] for v, _ in crows if v[0].startswith('11 kV protection relay'))
batt_src = next(v[1] for v, _ in crows if v[0].startswith('Battery system'))
RL, BT = '11 kV protection relay AH1 to AH7', 'Battery system ST3340UX'
insert_after(crows, 2, 'Disturbance record available', [
    [RL, relay_src, 'Trip circuit supervision alarm', 'Read', 'Status', '-', 7, 1, None, '-', REV],
    [RL, relay_src, 'Sensitive earth fault trip (51G)', 'Read', 'Event', '-', 6, 1, None, '-', REV + '. Panels with a zero-sequence CT (all except AH5)'],
    [RL, relay_src, 'Reverse power trip (32), AH1 generator incomer', 'Read', 'Event', '-', 1, 1, None, '-', REV + '. Protects the generators against reverse power from PV and battery'],
])
insert_after(crows, 2, 'Container fire status', [
    [BT, batt_src, 'Battery combiner panel DC switch closed (2 points)', 'Read', 'Status', '-', 1, 2, None, '-', REV + '. Combiner panels BCP unit 1# and 2#, if the EMS exposes them'],
    [BT, batt_src, 'Battery combiner panel fuse blown (2 points)', 'Read', 'Status', '-', 1, 2, None, '-', REV + '. As above'],
])
# redundant controller pair: health and supplies per controller, plus redundancy status
CPDEV = 'SICAM CP-8050 internal'
cp_src = next(v[1] for v, _ in crows if v[0] == CPDEV)
PAIR = ' (controllers A and B)'
for v, _ in crows:
    if v[0] == CPDEV and v[2] in ('Controller healthy', 'Power supply 1 healthy', 'Power supply 2 healthy'):
        v[2] += PAIR
        v[6] = 2
        v[10] = REV + '. One per controller of the hot standby pair'
    if v[0] == CPDEV and v[2] == 'I/O row 1 healthy':
        v[2] = 'I/O row 1 link healthy'
        v[10] = REV + '. Row 1 now has its own CI-8532 interface, so both controllers reach it'
insert_after(crows, 2, 'Link to each communication device healthy (8 links)', [
    [CPDEV, cp_src, 'Controller active' + PAIR, 'Read', 'Status', '-', 2, 1, None, '-', REV + '. Shows which controller is in control'],
    [CPDEV, cp_src, 'Standby controller synchronised and ready', 'Read', 'Status', '-', 1, 1, None, '-', REV + '. Off means a failure of the active controller would not be covered'],
    [CPDEV, cp_src, 'Controller switchover occurred', 'Read', 'Event', '-', 1, 1, None, '-', REV + '. Logged with time; resets on acknowledgement'],
])
for i, (v, _) in enumerate(crows):
    v[8] = '=G%d*H%d' % (i + 2, i + 2)
cp_new_last = 1 + len(crows)
for v, _ in foot:
    for k, x in enumerate(v):
        if isinstance(x, str) and x.startswith('='):
            v[k] = x.replace(str(cp_last), str(cp_new_last))
write_rows(cp, 2, crows + foot, cp.max_row + 10)

# ---------------------------------------------------------------- Device List
dl = wb['Device List']
dl_last = dl.max_row
drows = read_rows(dl, 2, dl_last)
insert_after(drows, 0, 'D17', [
    ['D49', 'Plant', 'LV main breakers: PV ACB 1250 A, BESS ACB 1000 A, DG step-up LV ACB; inverter AC MCCBs 400 A', 'TBC (switchboard supplier)', '-', '-',
     'Existing (drawings)', 3, 'Inverter MV station, BESS area, generator area', 'Dry contacts', 'Hardwired', 'Hardwired', 'Rows 5, 6',
     REV + '. Position and trip of each main breaker; trip of the four inverter MCCBs (row 5).'],
    ['D50', 'Plant', 'RMU UPS 2 kVA, 220 V AC', 'TBC (RMU supplier)', '-', '-', 'Existing (drawings)', 2, 'President and PM stations',
     'Dry contacts', 'Hardwired', 'Hardwired', 'Rows 3, 4', REV + '. On battery and fault contacts.'],
])
for v, _ in drows:
    if v[0] == 'D07':
        v[13] = (v[13] + ' ' if v[13] else '') + 'Breaker position, trip, earth switch and gas pressure also wired to row 1 (rev 1).'
    if v[0] == 'D04':
        v[13] = (v[13] + ' ' if v[13] else '') + 'Battery combiner panel status read through the EMS if available (rev 1).'
    if v[0] == 'D18':
        v[7] = 2
        v[8] = '11 kV switchroom, SCADA cabinet'
        v[13] = 'Hot standby pair (rev 1): controller A active, controller B standby, both on both core switches. Both connect to the SCADA servers by IEC 104. Siemens to confirm the redundancy function, switchover time and how the I/O rows connect to both controllers.'
    if v[0] == 'D19':
        v[7] = 4
        v[13] = 'Two per controller (rev 1). Confirm the CP-8050 supports a second supply.'
    if v[0] in ('D20', 'D21'):
        v[7] = 2
        v[13] = (v[13] + ' ' if v[13] else '') + 'One per controller (rev 1).'
    if v[0] == 'D24':
        v[13] = 'Alarm beacon and controller healthy output, driven by the active controller; the healthy output drops only if both controllers fail (rev 1).'
    if v[0] == 'D26':
        v[7] = 2
        v[8] = 'Rows 1, 2'
        v[12] = 'Rows 1, 2'
        v[13] = 'Row 1 moved off the controller backplane onto its own CI-8532 so both controllers reach it (rev 1). 24 V variant: CI-8530, 6MF2853-0AA00.'
    if v[0] == 'D28':
        v[7] = 2
        v[13] = 'One set per controller, including the redundancy function if licensed separately (rev 1).'
    if v[0] == 'D09':
        v[13] = (v[13] + ' ' if v[13] else '') + 'Auxiliary transformer also has its XRNT-3A fuse striker wired to row 2 (rev 1).'
write_rows(dl, 2, drows, dl_last + 2)
replace_in_formulas(wb['Protocol Summary'], '$%d' % dl_last, '$%d' % (dl_last + 2))

# ---------------------------------------------------------------- Comms Point List (expanded, tagged)
INST = {
    'PV inverter NB01 to NB04': ['NB01', 'NB02', 'NB03', 'NB04'],
    'Embedded EMS': ['EMS'], 'Battery system ST3340UX': ['BESS'],
    'Generator controller DG1, DG2': ['DG1', 'DG2'],
    '11 kV protection relay AH1 to AH7': ['AH1', 'AH2', 'AH3', 'AH4', 'AH5', 'AH6', 'AH7'],
    'RMU relay ST260E': ['PRS-RLY', 'PMS-RLY'],
    'Power meter PM5560': ['PRS-LV', 'PMS-LV', 'AUXLV'],
    "Weather logger blue'Log XM": ['WS'], 'UPS': ['UPS'],
    'Network switches (core and field)': ['SW-CORE-A', 'SW-CORE-B', 'SW-MV', 'SW-BESS', 'SW-GEN', 'SW-PRS', 'SW-PMS'],
    'Firewall and GPS time server': ['FW', 'GPS'], 'PLC box': ['PLCBOX'], 'SICAM CP-8050 internal': ['CP-A', 'CP-B', 'CP'],
}
PH3 = ['L1', 'L2', 'L3']
CODE = {
    'Active power': 'P', 'Reactive power': 'Q', 'Apparent power': 'S', 'Power factor': 'PF', 'Frequency': 'F',
    'AC voltage L1, L2, L3 (3 points)': ['U-' + p for p in PH3], 'AC current L1, L2, L3 (3 points)': ['I-' + p for p in PH3],
    'Voltage L1, L2, L3 (3 points)': ['U-' + p for p in PH3], 'Current L1, L2, L3 (3 points)': ['I-' + p for p in PH3],
    'Phase current L1, L2, L3 (3 points)': ['I-' + p for p in PH3], 'Voltage L-N, 3 phases (3 points)': ['U-' + p + 'N' for p in PH3],
    'Current, 3 phases (3 points)': ['I-' + p for p in PH3], 'Earth current': 'I-E',
    'DC voltage': 'U-DC', 'DC current': 'I-DC', 'DC power': 'P-DC', 'Energy today': 'E-DAY', 'Energy total': 'E-TOT',
    'Internal temperature': 'T-INT', 'Insulation resistance': 'R-ISO', 'Operating state': 'STATE', 'Fault code': 'FLT-CODE',
    'Alarm code': 'ALM-CODE', 'Communication healthy': 'COMM-OK', 'Active power limit': 'P-LIM-SP', 'Start / stop': 'START-STOP-CMD',
    'Plant active power': 'P-PLANT', 'Plant reactive power': 'Q-PLANT', 'PV total power': 'P-PV', 'BESS power': 'P-BESS',
    'Generator total power': 'P-GEN', 'Load power': 'P-LOAD', 'Bus voltage': 'U-BUS', 'Bus frequency': 'F-BUS',
    'Operating mode': 'MODE', 'EMS state': 'STATE', 'EMS summary alarm': 'ALM', 'Generator start request': 'GEN-START-REQ',
    'PV curtailment active': 'PV-CURT', 'EMS heartbeat': 'HEARTBEAT', 'Operating mode select': 'MODE-CMD',
    'State of charge minimum': 'SOC-MIN-SP', 'State of charge maximum': 'SOC-MAX-SP', 'Generator start threshold': 'GEN-START-SP',
    'Generator stop threshold': 'GEN-STOP-SP', 'BESS power limit': 'P-BESS-LIM-SP', 'BESS C-rate limit enable / disable': 'CRATE-CMD',
    'Fault reset, level 3': 'RESET-L3-CMD', 'Fault reset, levels 1 and 2': 'RESET-L12-CMD',
    'State of charge': 'SOC', 'State of health': 'SOH', 'Cell voltage maximum': 'U-CELL-MAX', 'Cell voltage minimum': 'U-CELL-MIN',
    'Cell temperature maximum': 'T-CELL-MAX', 'Cell temperature minimum': 'T-CELL-MIN', 'Charge power limit': 'P-CHG-LIM',
    'Discharge power limit': 'P-DIS-LIM', 'Energy charged total': 'E-CHG', 'Energy discharged total': 'E-DIS',
    'Converter state': 'PCS-STATE', 'Battery management state': 'BMS-STATE', 'Racks online': 'RACKS-ON', 'DC contactors closed': 'DC-CONT',
    'Cooling system state': 'COOL-STATE', 'Alarm level 1': 'ALM-L1', 'Alarm level 2': 'ALM-L2', 'Alarm level 3': 'ALM-L3',
    'Container fire status': 'FIRE', 'Battery combiner panel DC switch closed (2 points)': ['BCP1-SW-CLSD', 'BCP2-SW-CLSD'],
    'Battery combiner panel fuse blown (2 points)': ['BCP1-FUSE', 'BCP2-FUSE'],
    'Oil pressure': 'OIL-P', 'Coolant temperature': 'COOL-T', 'Starter battery voltage': 'BATT-U', 'Fuel level': 'FUEL',
    'Running hours': 'HOURS', 'Number of starts': 'STARTS', 'Controller mode': 'MODE', 'Breaker status': 'GCB', 'Warning': 'WARN',
    'Shutdown alarm': 'SHUTDOWN', 'Start': 'START-CMD', 'Stop': 'STOP-CMD', 'Mode auto / manual': 'MODE-CMD',
    'Energy import': 'E-IMP', 'Energy export': 'E-EXP', 'Breaker position': 'CB-POS', 'Overcurrent trip (50/51)': 'TRIP-50',
    'Earth fault trip (50N/51N)': 'TRIP-50N', 'Under / overvoltage trip (27/59)': 'TRIP-27-59', 'Frequency trip (81)': 'TRIP-81',
    'Relay healthy': 'RLY-HEALTHY', 'Disturbance record available': 'DREC', 'Trip circuit supervision alarm': 'TCS-ALM',
    'Sensitive earth fault trip (51G)': 'TRIP-51G', 'Reverse power trip (32), AH1 generator incomer': 'TRIP-32',
    'Overcurrent trip': 'TRIP-50', 'Earth fault trip': 'TRIP-50N', 'Active energy': 'E-P', 'Reactive energy': 'E-Q',
    'Voltage THD': 'THD-U', 'Current THD': 'THD-I', 'Plane-of-array irradiance': 'IRR-POA', 'Wind speed': 'WIND-SPD',
    'Wind direction': 'WIND-DIR', 'Ambient temperature': 'T-AMB', 'Module temperature': 'T-MOD', 'Logger healthy': 'HEALTHY',
    'Input voltage': 'U-IN', 'Output load': 'LOAD', 'Battery charge': 'BATT-CHG', 'Runtime remaining': 'RUNTIME',
    'Time synchronised / VPN session active': 'TIME-SYNC', 'On battery': 'ON-BATT', 'Fault': 'FLT', 'Device healthy': 'HEALTHY', 'Ring status': 'RING-OK',
    'Communication healthy per inverter (4 points)': ['NB01-COMM', 'NB02-COMM', 'NB03-COMM', 'NB04-COMM'],
    'Controller healthy (controllers A and B)': 'HEALTHY', 'Power supply 1 healthy (controllers A and B)': 'PS1-OK',
    'Power supply 2 healthy (controllers A and B)': 'PS2-OK', 'Controller active (controllers A and B)': 'ACTIVE',
    'Standby controller synchronised and ready': 'REDUN-OK', 'Controller switchover occurred': 'SWITCHOVER',
    'I/O row 1 link healthy': 'ROW1-OK', 'I/O row 2 link healthy': 'ROW2-OK', 'I/O row 3 link healthy': 'ROW3-OK',
    'I/O row 4 link healthy': 'ROW4-OK', 'I/O row 5 link healthy': 'ROW5-OK', 'I/O row 6 link healthy': 'ROW6-OK',
    'Link to each communication device healthy (8 links)': ['LINK%d-OK' % i for i in range(1, 9)],
}
# (level, alarm state) for status points; measured values carry no alarm here
ALARM = [
    (r'TRIP-|SHUTDOWN|^FIRE$|ALM-L1|BCP\d-FUSE', '1', 'On'), (r'RLY-HEALTHY|^HEALTHY$|COMM-OK|RING-OK|PS\d-OK|ROW\d-OK|LINK\d-OK|BCP\d-SW-CLSD|REDUN-OK', '2', 'Off'),
    (r'TCS-ALM|ALM-L2|^ALM$|^FLT$|ON-BATT', '2', 'On'), (r'WARN|ALM-L3|PV-CURT|GEN-START-REQ|DREC|VPN-ACTIVE|SWITCHOVER', '3', 'On'),
    (r'TIME-SYNC', '3', 'Off'),
]
if 'Comms Point List' in wb.sheetnames:
    del wb['Comms Point List']
ws = wb.create_sheet('Comms Point List', wb.sheetnames.index('Comms Points') + 1)
head = ['Tag', 'Device', 'Point', 'Direction', 'Data type', 'Unit', 'Source / protocol', 'Alarm level', 'Alarm when', 'Operated by', 'Notes']
ws.append(head)
hstyle = [copy(io.cell(row=1, column=c)._style) for c in range(1, 12)]
bstyle = copy(io.cell(row=2, column=1)._style)
out = []
devorder = []


def cp_point(point, inst):
    """'Controller healthy (controllers A and B)' -> 'Controller A healthy' for the tag of controller A."""
    if inst not in ('CP-A', 'CP-B') or not point.endswith(PAIR):
        return point
    base, c = point[:-len(PAIR)], inst[-1]
    return base.replace('Controller', 'Controller ' + c, 1) if base.startswith('Controller') else 'Controller ' + c + ' ' + base[0].lower() + base[1:]


for li, (v, _) in enumerate(crows):
    if v[0] not in devorder:
        devorder.append(v[0])
    dev, src, point, dirn, dtype, unit, ndev, each = v[0], v[1], v[2], v[3], v[4], v[5], v[6], v[7]
    insts = INST[dev]
    if point == 'Sensitive earth fault trip (51G)':
        insts = ['AH1', 'AH2', 'AH3', 'AH4', 'AH6', 'AH7']
    if point.startswith('Reverse power trip'):
        insts = ['AH1']
    if dev == CPDEV:
        insts = ['CP-A', 'CP-B'] if point.endswith(PAIR) else ['CP']
    assert len(insts) == ndev, (dev, point, insts, ndev)
    code = CODE[point]
    codes = code if isinstance(code, list) else [code]
    assert len(codes) == each, (point, codes, each)
    for ii, inst in enumerate(INST[v[0]] if v[0] in INST else insts):
        if inst not in insts:
            continue
        for ci, c in enumerate(codes):
            cc = c
            if dev.startswith('Firewall') and point.startswith('Time synchronised'):
                cc = 'TIME-SYNC' if inst == 'GPS' else 'VPN-ACTIVE'
            if inst.endswith('-RLY') and cc.startswith('RLY-'):
                cc = cc[4:]
            tag = inst + '-' + cc
            note = v[10] if v[10] else None
            if tag in tags:
                note = 'Communicated copy of hardwired signal %s' % tag
                tag += '-COM'
            lvl, when = ('Status', '-') if dtype in ('Status', 'Double point', 'Event') else ('-', '-')
            if dirn != 'Write':
                for pat, l, w in ALARM:
                    if re.search(pat, cc):
                        lvl, when = l, w
                        break
            out.append([(devorder.index(dev), ii, li, ci), tag, dev.replace('NB01 to NB04', inst).replace('DG1, DG2', inst).replace('AH1 to AH7', inst) if not (inst.startswith('CP') or inst == 'PLCBOX') else (dev.replace(' internal', ' ' + inst[-1]) if inst in ('CP-A', 'CP-B') else dev),
                        cp_point(point, inst), dirn, dtype, unit, src, lvl, when, v[9], note])
out = [o[1:] for o in sorted(out, key=lambda o: o[0])]
tl = [o[0] for o in out]
assert len(tl) == len(set(tl)), 'duplicate comms tag'
hw = set(tags)
assert not hw & set(tl), 'comms tag clashes with hardwired tag: %s' % (hw & set(tl))
for o in out:
    ws.append(o)
for c in range(1, 12):
    ws.cell(row=1, column=c)._style = copy(hstyle[c - 1])
for r in range(2, len(out) + 2):
    for c in range(1, 12):
        ws.cell(row=r, column=c)._style = copy(bstyle)
for col, w in zip('ABCDEFGHIJK', [22, 34, 44, 10, 12, 8, 30, 11, 11, 11, 50]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B2'
ws.auto_filter.ref = 'A1:K%d' % (len(out) + 1)

# ---------------------------------------------------------------- IO Rows: row 1 off the backplane
ior = wb['IO Rows']
for r in range(2, ior.max_row + 1):
    if ior.cell(row=r, column=1).value == 1:
        ior.cell(row=r, column=2).value = '11 kV switchroom, SCADA cabinet'
        ior.cell(row=r, column=3).value = 'CI-8532 (electrical)'
        ior.cell(row=r, column=4).value = '6MF2853-2AA00'
        ior.cell(row=r, column=5).value = 'Copper, inside the cabinet'
    for c in range(1, ior.max_column + 1):
        x = ior.cell(row=r, column=c).value
        if isinstance(x, str) and x.startswith('Each remote row is a small panel'):
            ior.cell(row=r, column=c).value = ('Each row has its own interface module and is reached by both controllers of the hot standby pair; '
                                               'only the active controller reads inputs and drives outputs. Each remote row is a small panel near its '
                                               'equipment, with field contacts wired a short distance to that panel.')
        if isinstance(x, str) and x.startswith('Confirm with Siemens'):
            ior.cell(row=r, column=c).value = x + ' Also confirm how the I/O rows connect to both controllers of the redundant pair.'

# ---------------------------------------------------------------- Datasheets
dsh = wb['Datasheets']
for r in range(2, dsh.max_row + 1):
    if dsh.cell(row=r, column=6).value == '6MF2805-0AA00':
        dsh.cell(row=r, column=7).value = 'Datasheet; CP-8050 system manual; redundancy description; security manual'

# ---------------------------------------------------------------- Read Me
rm = wb['Read Me']
for row in rm.iter_rows():
    for c in row:
        if c.value == 'Status':
            rm.cell(row=c.row, column=c.column + 1).value = (
                'Rev 1, 15 Oct 2026: harmonised with the design report, the drawings and the HMI demonstration. '
                'For client submission as Attachment B.')
        if c.value == 'Controller':
            rm.cell(row=c.row, column=c.column + 1).value = (
                'Two Siemens SICAM A8000 CP-8050 in hot standby (A active, B standby) as plant controller and gateway. '
                'The embedded EMS keeps dispatch of PV, battery and generators.')
        if c.value == 'SCADA':
            rm.cell(row=c.row, column=c.column + 1).value = (
                'Siemens SICAM SCC on two servers (active and hot standby), linked to both controllers by IEC 60870-5-104.')
        if c.value == 'Comms Points':
            rm.cell(row=c.row, column=c.column + 1).value = (
                'Typical signals read or written over communication links, per device type, with quantities. Replace with each vendor\'s point table.')
ins_at = next(c.row for row in rm.iter_rows() for c in row if c.value == 'Comms Points') + 1
rm.insert_rows(ins_at)
rm.cell(row=ins_at, column=1).value = 'Comms Point List'
rm.cell(row=ins_at, column=2).value = 'The Comms Points expanded to one row per signal, each with its own tag and alarm level, as used by the HMI.'
for c in (1, 2):
    rm.cell(row=ins_at, column=c)._style = copy(rm.cell(row=ins_at - 1, column=c)._style)
nr = rm.max_row + 2
rm.cell(row=nr, column=1).value = 'Revision 1 changes'
rm.cell(row=nr, column=1)._style = copy(rm.cell(row=ins_at - 1, column=1)._style)
changes = [
    'AH8: breaker open, closed, tripped, earth switch and gas pressure added (row 1). The 11 kV system is earthed only through AH8.',
    'LV main breakers added: PV ACB 1250 A and BESS ACB 1000 A (row 5), DG step-up transformer LV ACB (row 6).',
    'Combiner panels: trip of the four inverter AC MCCBs (row 5); battery combiner panel switch and fuse status through the EMS.',
    'Auxiliary transformer XRNT-3A fuse blown (row 2); RMU UPS 2 kVA on battery and fault at both office stations (rows 3, 4).',
    '11 kV relays: trip circuit supervision (7), sensitive earth fault trip (6) and reverse power trip on AH1 added as communication points.',
    'Totals: 244 hardwired signals and %d communication points.' % len(out) + ' Overall spare is 27%; row 5 drops to 19% (6 of 32 channels), just under the 20% target: add a third DI-8112 there if more yard signals appear.',
    'Comms Point List sheet added. Devices D49 and D50 added to the Device List.',
    'Controller made redundant: two CP-8050 (A active, B hot standby), each with two power supplies, its own Ethernet interface, flash card '
    'and licences. I/O row 1 moved off the controller backplane onto its own CI-8532 so both controllers reach it. Communication points added '
    'for each controller\'s health and supplies, which controller is active, standby ready and switchover. Siemens to confirm the redundancy '
    'function, switchover time and how the I/O rows connect to both controllers.',
]
for i, t in enumerate(changes):
    rm.cell(row=nr + 1 + i, column=1).value = i + 1
    rm.cell(row=nr + 1 + i, column=2).value = t

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('IO List rows', len(rows), 'comms lines', len(crows), 'comms points expanded', len(out))
