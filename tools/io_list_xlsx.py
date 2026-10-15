"""Build Annex A (I/O list) as a formatted Excel workbook from docs/IO_List.csv.
Usage: python3 tools/io_list_xlsx.py   (run tools/export-io-list.js first)"""
import csv, os
from collections import Counter, OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(__file__), '..')
src = os.path.join(ROOT, 'docs', 'IO_List.csv')
out = os.path.join(ROOT, 'docs', 'Annex_A_IO_List.xlsx')

with open(src, encoding='utf-8-sig') as f:
    rows = list(csv.reader(f))
head, data = rows[0], rows[1:]

navy = PatternFill('solid', fgColor='0E2F47')
band = PatternFill('solid', fgColor='EEF3F7')
hfont = Font(bold=True, color='FFFFFF')
thin = Side(style='thin', color='C9D2DA')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# --- cover / summary sheet
s = wb.active
s.title = 'Summary'
s['A1'] = 'Villa Somalia Hybrid Power Plant - SCADA I/O List (Annex A)'
s['A1'].font = Font(bold=True, size=14)
s['A2'] = 'Ref. BFE/PR/26/01-SC-001 Rev 0 - Issued for approval'
s['A3'] = 'Generated from the SCADA tag database that drives the HMI screens. Modbus addresses are placeholders until vendor register maps are received.'
s['A3'].alignment = Alignment(wrap_text=True)
s.merge_cells('A3:G3'); s.row_dimensions[3].height = 32
types = ['AI', 'DI', 'AO', 'DO']
labels = {'AI': 'Analogue in', 'DI': 'Digital in', 'AO': 'Setpoints (AO)', 'DO': 'Commands (DO)'}
cnt = OrderedDict()
for r in data:
    cnt.setdefault(r[1], Counter())[r[4]] += 1
hdr = ['Area'] + [labels[t] for t in types] + ['Total']
for i, h in enumerate(hdr, 1):
    c = s.cell(row=5, column=i, value=h); c.fill = navy; c.font = hfont; c.border = box
r_ = 6
for area, c in cnt.items():
    vals = [area] + [c[t] for t in types] + [sum(c.values())]
    for i, v in enumerate(vals, 1):
        cell = s.cell(row=r_, column=i, value=v); cell.border = box
    r_ += 1
tot = ['Total'] + [sum(c[t] for c in cnt.values()) for t in types] + [len(data)]
for i, v in enumerate(tot, 1):
    cell = s.cell(row=r_, column=i, value=v); cell.font = Font(bold=True); cell.border = box
s.column_dimensions['A'].width = 38
for col in 'BCDEF':
    s.column_dimensions[col].width = 15
s['A' + str(r_ + 2)] = 'Alarm priority: 1 critical (trips, fire, loss of supply) - 2 high - 3 medium - 4 low.'
s['A' + str(r_ + 3)] = 'Tag format: AREA.EQUIPMENT.SIGNAL'

# --- full list
w = wb.create_sheet('I-O List')
for i, h in enumerate(head, 1):
    c = w.cell(row=1, column=i, value=h); c.fill = navy; c.font = hfont
    c.alignment = Alignment(vertical='center', wrap_text=True); c.border = box
numcols = {head.index(x) for x in ('Priority', 'LoLo', 'Lo', 'Hi', 'HiHi')}
for ri, r in enumerate(data, 2):
    for ci, v in enumerate(r):
        if ci in numcols and v != '':
            try: v = float(v) if '.' in v else int(v)
            except ValueError: pass
        c = w.cell(row=ri, column=ci + 1, value=v); c.border = box
        if ri % 2 == 0: c.fill = band
widths = [30, 26, 16, 16, 6, 48, 8, 12, 16, 34, 30, 14, 8, 7, 7, 7, 7]
for i, wd in enumerate(widths, 1):
    w.column_dimensions[get_column_letter(i)].width = wd
w.freeze_panes = 'B2'
w.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(head)), len(data) + 1)
w.row_dimensions[1].height = 30
wb.save(out)
print('wrote', out, len(data), 'points')
