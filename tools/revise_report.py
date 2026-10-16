"""Rev 1 of the Villa Somalia SCADA Design Report: harmonised with I/O list rev 1.

Usage: python3 tools/revise_report.py <rev0.docx> <out.docx>
Each edit names the exact old text, so a changed source fails loudly instead of
editing the wrong place.
"""
import os
import sys
from copy import deepcopy
import docx

SRC, OUT = sys.argv[1], sys.argv[2]
COMMS = 463  # communication points in I/O list rev 1 (tools/harmonize_io_list.py)
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'docs', 'figures')
d = docx.Document(SRC)
T = d.tables


def set_par(p, old, new):
    """Replace old with new inside one paragraph, keeping the run formatting around it."""
    text = ''.join(r.text for r in p.runs)
    assert old in text, (old, text)
    start = text.index(old)
    pos = 0
    first = None
    for r in p.runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= start or a >= start + len(old):
            continue
        if first is None:
            first = r
            r.text = r.text[:start - a] + new + (r.text[start + len(old) - a:] if b >= start + len(old) else '')
        else:
            r.text = r.text[start + len(old) - a:] if b > start + len(old) else ''


def cell(t, r, c, old, new):
    ps = [p for p in T[t].rows[r].cells[c].paragraphs if old in p.text]
    assert ps, ('T%d R%d C%d' % (t, r, c), old, T[t].rows[r].cells[c].text)
    set_par(ps[0], old, new)


def par(old, new):
    ps = [p for p in d.paragraphs if old in p.text]
    assert len(ps) == 1, (old, len(ps))
    set_par(ps[0], old, new)
    return ps[0]


def insert_par_after(p, text):
    el = deepcopy(p._p)
    p._p.addnext(el)
    np = docx.text.paragraph.Paragraph(el, p._parent)
    for r in np.runs[1:]:
        r._r.getparent().remove(r._r)
    np.runs[0].text = text
    return np


def add_row(t, values):
    tr = deepcopy(T[t].rows[-1]._tr)
    T[t]._tbl.append(tr)
    row = T[t].rows[-1]
    for c, v in zip(row.cells, values):
        p = c.paragraphs[0]
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = v


XLSX_OLD, XLSX_NEW = 'Villa_Somalia_SCADA_Device_and_IO_List.xlsx', 'Villa_Somalia_SCADA_Device_and_IO_List_Rev1.xlsx'

# 1 Reference documents and attachments
cell(0, 3, 0, XLSX_OLD, XLSX_NEW)
cell(28, 2, 1, XLSX_OLD, XLSX_NEW)
cell(28, 2, 2, 'Draft complete', 'Rev 1, complete')

# 2 Compliance table, client comment 3 (combiner panels now covered)
cell(1, 3, 2, '224 hardwired signals and about 440 communication points, listed signal by signal',
     '244 hardwired signals and ' + str(COMMS) + ' communication points, listed signal by signal, including the PV inverter AC breakers and the battery combiner panels')

# 3 Design basis: generator ratings from the SLD
cell(2, 6, 1, 'Ratings per generator supplier', 'DG1 550 kVA, DG2 660 kVA; step-up transformer 1.5 MVA, 11/0.4 kV')

# 5 Scope of monitoring and control
cell(7, 1, 2, 'None', 'Main LV breaker and the four inverter AC breakers to row 5')
cell(7, 3, 2, 'Fire, E-stop and door to row 5', 'Fire, E-stop, door and main LV breaker to row 5')
cell(7, 3, 3, 'charge and discharge limits, alarms', 'charge and discharge limits, combiner panel status, alarms')
cell(7, 4, 2, 'fuel level to row 6', 'fuel level and step-up transformer LV breaker to row 6')
cell(7, 5, 2, 'relay healthy to row 1', 'relay healthy to row 1, including the AH8 breaker')
cell(7, 6, 2, 'local/remote to rows 3 and 4', 'local/remote and RMU UPS to rows 3 and 4')
cell(7, 8, 2, 'Temperature, pressure and oil level contacts to the nearest row',
     'Temperature, pressure and oil level contacts, and the auxiliary transformer fuse, to the nearest row')

# 6 PV and BESS interface
cell(8, 3, 2, 'Battery fire panel, emergency stop and container door to I/O row 5',
     'Battery fire panel, emergency stop, container door and the PV and battery LV breakers to I/O row 5')
cell(8, 3, 5, 'Fire panel states, emergency stop, door open', 'Fire panel states, emergency stop, door open, LV breaker positions and trips')
cell(10, 3, 1, 'About 23', 'About 27')
cell(10, 3, 3, 'alarm levels', 'alarm levels, combiner panel status')
cell(10, 4, 1, '10', '18')
cell(10, 4, 3, 'Fire panel states, emergency stop, container door', 'Fire panel states, emergency stop, container door, LV breakers, inverter AC breakers')

# 7 I/O summary
par('The plant has 224 hardwired signals on 24 I/O modules in six rows, with 33% spare channels. A further 440 or so signals arrive over communication links. The full signal-by-signal list is in Attachment B.',
    'The plant has 244 hardwired signals on 24 I/O modules in six rows, with 27% spare channels. Row 5, in the BESS and inverter yard, is the fullest at 19% spare; a third input module can be added there if more yard signals appear. A further ' + str(COMMS) + ' signals arrive over communication links. The full signal-by-signal list is in Attachment B: sheet IO List for hardwired signals and sheet Comms Point List for communication points.')
for r, old, new in ((1, '82', '87'), (2, '24', '25'), (3, '41', '43'), (4, '41', '43'), (5, '18', '26'), (6, '18', '20'), (7, '224', '244')):
    cell(12, r, 6, old, new)
cell(12, 3, 4, 'RMU, 630 kVA transformer, LVDB', 'RMU, 630 kVA transformer, LVDB, RMU UPS')
cell(12, 4, 4, 'RMU, 630 kVA transformer, LVDB', 'RMU, 630 kVA transformer, LVDB, RMU UPS')
cell(12, 5, 4, 'Inverter transformer, fire panel, BESS E-stop and door', 'Inverter transformer, PV and BESS LV breakers, inverter AC breakers, fire panel, BESS E-stop and door')
cell(12, 6, 4, 'Generator status and fuel level, DG step-up transformer', 'Generator status and fuel level, DG step-up transformer and its LV breaker')
cell(13, 1, 4, '188', '208')
cell(13, 1, 5, '84', '64')
cell(13, 5, 4, '224', '244')
cell(13, 5, 5, '112', '92')
cell(14, 1, 1, '188', '208')
cell(14, 4, 1, 'About 440', str(COMMS))

# 10 Breaker switching rules: earthing interlock
par('such as closing a breaker onto an earthed circuit.', 'such as closing a breaker onto an earthed circuit, or energising the busbar through AH1 or AH3 while AH8 is open.')

# 11 Black start: earth the 11 kV system through AH8 first
p = par('After a total shutdown, all 11 kV breakers are open.',
        'After a total shutdown, all 11 kV breakers are open except AH8, which stays closed unless it has tripped.')
insert_par_after(p, 'AH8 is confirmed closed, or closed first, so the 11 kV system is earthed through the neutral earthing resistor before the busbar is energised. The controller blocks closing AH1 or AH3 while AH8 is open.')

# 13 24 V power supply: match the Device List and Datasheets register
cell(25, 12, 1, 'Phoenix Contact QUINT4-PS/1AC/24DC/20', 'Siemens SITOP PSU8200 24 V/20 A')
par('Phoenix Contact QUINT4', 'Siemens SITOP PSU8200')

# 14 Open points
add_row(27, ['14', 'Voltage transformers on the 11 kV panels or busbar. None are shown on the drawings, yet the relay voltages and the 27/59 and 81 protection need them',
             'Switchgear supplier and electrical designer', 'Sections 5, 7; relay communication points'])
add_row(27, ['15', 'Whether the EMS exposes the battery combiner panel status', 'PV and BESS supplier', 'Battery communication points'])

# ---- Rev 1: requirements on the EMS and fallback (generator path, PV curtailment, black start, EMS failure, load shedding)
def clone_after(src, anchor, lead, text=None):
    """Copy paragraph src (style, numbering, run format) after anchor, with a bold lead-in and plain text."""
    el = deepcopy(src._p)
    anchor._p.addnext(el)
    np = docx.text.paragraph.Paragraph(el, anchor._parent)
    runs = np.runs
    for r in runs[2 if text is not None else 1:]:
        r._r.getparent().remove(r._r)
    runs[0].text = lead
    if text is not None:
        runs[1].text = ' ' + text
    return np


def find(prefix):
    ps = [p for p in d.paragraphs if p.text.startswith(prefix)]
    assert len(ps) == 1, (prefix, len(ps))
    return ps[0]


H3, NORMAL, BULLET = find('Black start'), find('The SCADA shows each step and raises'), find('Audit log.')
p = par('The EMS supplier must confirm which of these steps it carries out automatically and which need an operator command from the SCADA.',
        'The EMS supplier provides this sequence in its functional description, states which steps it carries out automatically and which need an operator command from the SCADA, and demonstrates it at commissioning.')
a = clone_after(H3, p, 'Requirements on the EMS and fallback')
a = clone_after(NORMAL, a, 'The design relies on the embedded EMS for dispatch, with no separate plant controller. The EMS supplier must therefore provide the functions below, and the plant must keep supplying the offices if the EMS fails. They are part of the EMS specification and are tested at commissioning.')
for lead, text in (
    ('Generator start and stop.', 'The EMS starts and stops each generator through a hardwired start contact to the generator controller, so the start does not depend on the CP-8050, the SCADA or the network. A Modbus command may duplicate it. The EMS applies the start below 20% and stop at 80% state of charge, the second-generator rule at 80% of the running generator, the 30% minimum load and lead and lag alternation.'),
    ('PV curtailment.', 'The EMS limits the inverters through the PLC box when the battery is full, or when a generator would fall below its minimum load. Output is ramped down rather than tripped. The supplier states the response time from the EMS to the inverters.'),
    ('Black start.', 'The EMS supplier provides the black-start sequence above in writing, including AH8 being closed first, and marks each step as automatic or by operator command.'),
    ('EMS failure.', 'The CP-8050 monitors the EMS heartbeat and raises a level 1 alarm when it stops. Each generator controller is set to start on its own on a dead busbar, or when the EMS start contact is lost, and to form the grid alone. The offices are then resupplied from the generators without the EMS, and the battery stays off until the EMS is restored.'),
    ('Load shedding.', 'If a generator is overloaded and frequency falls, the 11 kV relays shed load by under-frequency protection (81): AH6, the Prime Minister feeder, first, and AH4, the President feeder, last. Proposed settings are 48.5 Hz for 1 s on AH6 and 47.5 Hz for 0.5 s on AH4, to be confirmed by the protection study. The auxiliary feeder AH5 is not shed. Feeders are reclosed by operator command once frequency is stable.'),
):
    a = clone_after(BULLET, a, lead, text)

# division of duties, control authority, fault responses, recovery
cell(9, 5, 1, 'EMS and generator controllers', 'EMS, by hardwired start contact to the generator controllers')
add_row(6, ['Generator start if the EMS fails', 'Generator controllers, on dead busbar', 'Continues'])
add_row(6, ['Load shedding on under-frequency', 'Relays AH6 and AH4', 'Continues'])
add_row(9, ['Load shedding on under-frequency', 'Relays AH6, then AH4', 'Protection engineer, locally', 'Monitor trips; reclose feeders'])
add_row(22, ['EMS failure or EMS heartbeat lost', 'Level 1 alarm; generators start on dead busbar and supply the plant; battery off', 'Call the PV and BESS supplier; check generator loading'])
add_row(22, ['Generator overload, frequency falling', 'Relays shed AH6, then AH4, by under-frequency protection', 'Reduce load; reclose the feeders from the SCADA when frequency is stable'])
par('A written procedure covers local operation of breakers and generators if the controller or SCADA fails.',
    'A written procedure covers local operation of breakers and generators if the controller or SCADA fails. If the EMS fails, the generators start on their own and supply the plant (section 11).')

# open points
add_row(27, ['16', 'EMS functional description: generator start contact and Modbus commands, PV curtailment path and response time, black-start steps marked automatic or by operator, behaviour on loss of the CP-8050 link', 'PV and BESS supplier', 'Section 11; EMS communication points'])
add_row(27, ['17', 'Generator controller start input for the EMS contact; dead-bus automatic start and island operation without the EMS', 'Generator supplier', 'Section 11; generator settings'])
add_row(27, ['18', 'Under-frequency load-shedding settings and shedding order (AH6 first, AH4 last)', 'Client and protection engineer', 'Section 11; relay settings'])

# ---- Rev 1: CP-8050 as a hot standby pair (controller A active, B standby)
par('One Siemens SICAM A8000 CP-8050 controller collects every signal in the plant and passes it to two SCADA servers.',
    'A hot standby pair of Siemens SICAM A8000 CP-8050 controllers collects every signal in the plant and passes it to two SCADA servers. Controller A is in control; controller B follows it and takes over if A fails.')
par('1 controller, 5 network sources', 'redundant controller pair, 5 network sources')
par('The SCADA servers, workstation, firewall, controller and 11 kV relays each connect to both core switches.',
    'The SCADA servers, workstation, firewall, both controllers and the 11 kV relays each connect to both core switches.')
par('one link carries all its signals to the controller.',
    'one link carries all its signals to the controllers. Both controllers reach every row; only the active one reads inputs and drives outputs.')
cell(5, 1, 0, 'SICAM A8000 CP-8050', 'SICAM A8000 CP-8050, hot standby pair A and B')
cell(5, 1, 1, 'passes settings to the EMS', 'passes settings to the EMS; B takes over if A fails')
cell(6, 4, 2, 'Lost until restored', 'Continue on the standby controller or server; lost only if both fail')
cell(6, 5, 2, 'Breakers operated locally at the panels', 'Continues on the standby; if both controllers fail, breakers are operated locally')
cell(12, 1, 1, 'beside the controller', 'SCADA cabinet')
cell(12, 1, 2, 'None', 'CI-8532, electrical')
cell(12, 1, 3, 'Backplane', 'Copper')
cell(15, 1, 0, 'Controller to SCADA servers A and B', 'Controllers A and B to SCADA servers A and B')
cell(15, 9, 0, 'Controller to I/O row 2', 'Controllers to I/O rows 1 and 2')
add_row(15, ['Controller A to controller B', 'Redundancy synchronisation (Siemens)', 'Direct cable between the controllers, to confirm', 'Through the core switches, to confirm'])
cell(16, 7, 0, 'Controller fails', 'Active controller fails')
cell(16, 7, 1, 'SCADA loses its view; the EMS, relays and hardwired trips continue; the cold spare is fitted',
     'The standby controller takes over; the SCADA view and remote switching continue; a level 2 alarm is raised')
add_row(16, ['Both controllers fail', 'SCADA loses its view; the EMS, relays and hardwired trips continue; breakers are operated locally'])
par('A cold spare is included; a hot standby pair is offered as an option.',
    'Controllers A and B run as a hot standby pair, so it is no longer a single point. Each I/O row interface remains single; its failure loses only that row.')
cell(22, 6, 0, 'Controller failure', 'Active controller failure')
cell(22, 6, 1, 'Watchdog contact drops; alarm beacon sounds; EMS and relays continue', 'Standby controller takes over; level 2 alarm; SCADA view and remote switching continue')
cell(22, 6, 2, 'Operate locally under the recovery procedure; fit the spare controller', 'Repair or replace the failed controller; it rejoins as standby')
add_row(22, ['Failure of both controllers', 'Watchdog contact drops; alarm beacon sounds; EMS and relays continue', 'Operate locally under the recovery procedure'])
par('if the controller or SCADA fails.', 'if both controllers or both SCADA servers fail.')
cell(24, 1, 0, 'Plant controller (processor module)', 'Plant controller (processor module), hot standby pair')
for r, old, new in ((1, '1', '2'), (2, '2', '4'), (3, '1', '2'), (4, '1', '2'), (9, '1', '2'), (11, '1 set', '2 sets')):
    cell(24, r, 3, old, new)
cell(24, 11, 0, 'IEC 61850', 'IEC 61850, redundancy if licensed separately')
cell(26, 1, 0, 'CP-8050 with configured flash card (cold spare)', 'Configured flash card for the CP-8050')
cell(26, 1, 2, 'Restores the controller in under an hour instead of waiting for delivery', 'Controller B is the running spare; the card restores a replaced controller quickly')
par('A hot standby controller pair is available as an option for the client to decide. It needs Siemens to confirm how the remote I/O rows connect to two controllers.',
    'The controllers are supplied as a hot standby pair. I/O row 1 has its own CI-8532 interface rather than sitting on the controller backplane, so both controllers reach it. Siemens will confirm the redundancy function and licence, the switchover time, and how the I/O rows connect to both controllers.')
cell(27, 12, 1, 'Hot standby controller pair as an option', 'Hot standby controller pair: redundancy function and licence, switchover time, connection of the I/O rows to both controllers, product numbers')
cell(27, 12, 2, 'Client decision', 'Siemens')
cell(27, 12, 3, 'Section 13', 'Sections 4, 7, 8, 13')

# figures 1 and 3 redrawn for the controller pair (tools/draw_report_figures.py)
from PIL import Image
for k, name in ((0, 'fig1_architecture.png'), (2, 'fig3_network.png')):
    shape = d.inline_shapes[k]
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    path = os.path.join(FIG, name)
    d.part.related_parts[rid]._blob = open(path, 'rb').read()
    w, h = Image.open(path).size
    shape.height = int(shape.width * h / w)

# ---- Rev 1: HMI screen mock-ups in section 12 and as Attachment E (tools/capture_screens.js, tools/build_attachment_e.js)
from docx.shared import Emu
SCR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'docs', 'Attachment_E_Screens')
par('Screens will be submitted as mock-ups for approval before configuration.',
    'Mock-ups of the main screens are shown below and the full set is in Attachment E. They are submitted for approval of layout and content; the final screens are built in SICAM SCC to the same layout before the factory test.')
add_row(23, ['I/O list', 'Every signal with its tag, I/O row, module and channel or protocol source, alarm level and live value', 'Filter, export'])
caption_src = find('SCADA architecture · ')
anchor = T[23]._tbl
for name, text in (('E01_plant_overview.png', 'Plant overview · mock-up of the opening screen, a live single-line diagram of the whole plant'),
                   ('E02_11kV_switchboard.png', '11 kV switchboard · mock-up, panels AH1 to AH8 with positions, relay status and measurements'),
                   ('E07_energy_management.png', 'Energy management · mock-up, EMS state, power balance, operating modes and Admin settings')):
    fp = d.add_paragraph()
    fp.add_run().add_picture(os.path.join(SCR, name), width=Emu(5486400))
    anchor.addnext(fp._p)
    cap = deepcopy(caption_src._p)
    fp._p.addnext(cap)
    cp = docx.text.paragraph.Paragraph(cap, caption_src._parent)
    for r in cp.runs[1:]:
        r._r.getparent().remove(r._r)
    cp.runs[0].text = text
    anchor = cap
add_row(28, ['E', 'HMI screen mock-ups: 12 screens from the HMI demonstration, including a battery fire (Villa_Somalia_SCADA_Attachment_E_HMI_Screens.pdf)', 'Rev 1, complete; final screens built in SICAM SCC'])
cell(1, 1, 2, 'attachments A to D', 'attachments A to E')
cell(1, 2, 2, 'One controller and one SCADA covering every equipment group', 'A redundant controller pair and one SCADA covering every equipment group, with screen mock-ups')
cell(1, 2, 3, '3, 4, 5, 6, 12', '3, 4, 5, 6, 12, Attachment E')

d.save(OUT)
print('saved', OUT)
