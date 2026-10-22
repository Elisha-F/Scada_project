"""Bill of material for the SICAM A8000 design (Rev 1 architecture, drawing set BFE/PR/26/01-SCADA-001 to 005).

Usage: python3 tools/build_bom_sicam.py [out.xlsx]
Quantities follow the Device and I/O List Rev 1 (docs/superseded); prices are entered once on the Summary sheet.
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else 'docs/Villa_Somalia_SCADA_Bill_of_Material_SICAM.xlsx'

NAVY, AMBER, GREY, YELLOW, BAND = '14263D', 'E8A33B', 'EEF1F4', 'FFF2CC', 'DCE6F0'
THIN = Side(style='thin', color='C9D0D8')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FONT = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
SEC_FILL = PatternFill('solid', fgColor=BAND)
TBC_FILL = PatternFill('solid', fgColor=YELLOW)
WRAP = Alignment(wrap_text=True, vertical='top')

CHK, MEM, SUP, TBC = 'Checked', 'From memory - confirm', 'To confirm with supplier', 'TBC'

# Catalogue: key -> (description, manufacturer, model, order no., number status, unit)
CAT = {
    'CP8050': ('Processor module, SICAM A8000', 'Siemens', 'CP-8050', '6MF2805-0AA00', CHK, 'pc'),
    'PS8642': ('Power supply 110-220 V DC, 45 W', 'Siemens', 'PS-8642', '6MF2864-2AA00', CHK, 'pc'),
    'CI8520': ('Ethernet interface, 5 x RJ45', 'Siemens', 'CI-8520', '6MF2852-0AA00', CHK, 'pc'),
    'FLASH': ('SICAM flash card, 2 GB', 'Siemens', 'SICAM flash card', '6MF1213-1GA05-0AA0', CHK, 'pc'),
    'LIC': ('Protocol licences for one CP-8050: IEC 60870-5-104 server, Modbus TCP client, IEC 61850 client, hot standby function',
            'Siemens', 'CP-8050 licences', 'By quotation', SUP, 'set'),
    'CI8532': ('I/O remote interface, electrical, 110-220 V DC', 'Siemens', 'CI-8532', '6MF2853-2AA00', CHK, 'pc'),
    'CI8533': ('I/O remote interface, optical, 110-220 V DC', 'Siemens', 'CI-8533', 'To confirm', SUP, 'pc'),
    'DI8112': ('Digital input, 2 x 8 inputs, 110 V DC', 'Siemens', 'DI-8112', '6MF2811-2AA00', CHK, 'pc'),
    'DO8221': ('Secured command output, 12 outputs, 1-out-of-n check', 'Siemens', 'DO-8221', 'To confirm', SUP, 'pc'),
    'DO8212': ('Digital output, 8 relays, 24-220 V DC / 230 V AC', 'Siemens', 'DO-8212', '6MF2821-2AA00', CHK, 'pc'),
    'AI8320': ('Analogue input, 4 x 4-20 mA', 'Siemens', 'AI-8320', '6MF2832-0AA00', CHK, 'pc'),
    'CORE': ('Managed core switch, 8 x RJ45 + 4 x gigabit SFP', 'Moxa', 'EDS-G512E', 'EDS-G512E-4GSFP', MEM, 'pc'),
    'FSW': ('Managed field switch, 7 x RJ45 + 3 x gigabit combo/SFP', 'Moxa', 'EDS-510E', 'EDS-510E-3GTXSFP', MEM, 'pc'),
    'SFP': ('SFP transceiver, gigabit, single-mode, LC', 'Moxa', 'SFP-1GLXLC', 'SFP-1GLXLC', MEM, 'pc'),
    'MGATE': ('Modbus RTU to Modbus TCP gateway, 1 port', 'Moxa', 'MGate MB3170', 'MB3170', MEM, 'pc'),
    'FW': ('Rugged firewall with VPN', 'Fortinet', 'FortiGate Rugged 60F', 'FGR-60F', MEM, 'pc'),
    'GPS': ('GPS time server with antenna, cable and surge protector', 'Meinberg', 'LANTIME M300', 'M300/GPS', MEM, 'set'),
    'PSU20': ('24 V DC power supply 20 A, input 85-275 V AC / 88-350 V DC', 'Siemens', 'SITOP PSU8200 24 V/20 A', '6EP1336-3BA10', CHK, 'pc'),
    'PSU5': ('24 V DC power supply 5 A, input 85-264 V AC / 110-300 V DC', 'Siemens', 'SITOP PSU8200 24 V/5 A', '6EP1333-3BA10', MEM, 'pc'),
    'PM5560': ('Power meter with Ethernet, Modbus TCP', 'Schneider Electric', 'PowerLogic PM5560', 'METSEPM5560', MEM, 'pc'),
    'VX25': ('SCADA cabinet 800 x 2000 x 800 mm, with plinth', 'Rittal', 'VX25', '8808.000', MEM, 'pc'),
    'AX': ('Outdoor I/O panel, IP55 minimum, with sun shade', 'Rittal', 'AX', 'Size to suit', TBC, 'pc'),
    'AXS': ('Small wall enclosure for field switch, IP55', 'Rittal', 'AX', 'Size to suit', TBC, 'pc'),
    'ACC_C': ('Cabinet accessories: MCBs, terminals, surge protection, fans and thermostat, lighting, door switch, labels, wiring',
              'To suit', '-', '-', TBC, 'lot'),
    'ACC_P': ('Panel accessories: MCBs, terminals, surge protection, anti-condensation heater and thermostat, door switch, labels, wiring',
              'To suit', '-', '-', TBC, 'lot'),
    'FPP': ('Fibre patch panel, 24 x LC, rack or DIN', 'To suit', '-', '-', TBC, 'pc'),
    'FSB': ('Fibre splice box, DIN rail, 12 x LC', 'To suit', '-', '-', TBC, 'pc'),
    'IPC': ('SCADA server, rack industrial PC, redundant power supplies and disks', 'Siemens', 'SIMATIC IPC647E', 'Configured to order', SUP, 'pc'),
    'RACK': ('19-inch rack 24U with rack console (monitor, keyboard, KVM)', 'To suit', '-', '-', TBC, 'set'),
    'OWS': ('Operator workstation', 'Dell', 'OptiPlex 7000 series', 'Configured to order', SUP, 'pc'),
    'MON': ('Monitor 24-inch, full HD', 'Dell', 'P2425H or equal', 'To suit', SUP, 'pc'),
    'LAP': ('Engineering laptop, rugged', 'Dell', 'Latitude 5430 Rugged', 'Configured to order', SUP, 'pc'),
    'UPS': ('UPS 3000 VA online, rack, network card', 'APC', 'Smart-UPS SRT 3000', 'SRT3000RMXLI', MEM, 'pc'),
    'UPSB': ('External battery pack for the UPS (about 2 h runtime, to size)', 'APC', 'SRT96RMBP', 'SRT96RMBP', MEM, 'pc'),
    'BEACON': ('Alarm sounder and beacon', 'Pfannenberg', 'PATROL', 'Model to suit', TBC, 'pc'),
    'FIBRE': ('Single-mode OS2 fibre cable, 12 core, armoured, outdoor', 'To suit', '-', '-', TBC, 'm'),
    'FCORD': ('Fibre patch cord LC-LC single-mode duplex, 2 m', 'To suit', '-', '-', TBC, 'pc'),
    'CAT6': ('Cat6 S/FTP industrial cable', 'To suit', '-', '-', TBC, 'm'),
    'PCORD': ('Cat6 S/FTP patch cord, 1 to 3 m', 'To suit', '-', '-', TBC, 'pc'),
    'RJ45': ('Shielded field RJ45 connector', 'To suit', '-', '-', TBC, 'pc'),
    'RS485': ('RS485 cable, 2-pair shielded twisted pair, 120 ohm', 'To suit', '-', '-', TBC, 'm'),
    'CTRL': ('Control cable, multicore 1.0 mm2, screened, for hardwired I/O', 'To suit', '-', '-', TBC, 'm'),
    'SCC': ('Substation SCADA, redundant server pair, about 1,000 data points, IEC 60870-5-104, two user roles', 'Siemens', 'SICAM SCC',
            'By quotation', SUP, 'lic'),
    'SCCC': ('SICAM SCC client licence for the operator workstation', 'Siemens', 'SICAM SCC client', 'By quotation', SUP, 'lic'),
    'DM': ('Controller engineering tool', 'Siemens', 'SICAM Device Manager', 'By quotation', SUP, 'lic'),
    'WINS': ('Server operating system', 'Microsoft', 'Windows Server Standard', 'OEM with the IPC', SUP, 'lic'),
    'NMS': ('Network management, at least 16 nodes', 'Moxa', 'MXview One', 'By quotation', SUP, 'lic'),
    'FWLIC': ('Firewall security services and support, 3 years', 'Fortinet', 'FortiGate Rugged 60F bundle', 'By quotation', SUP, 'lic'),
}

# BOM by panel: (group code, group title, location, [(item key, qty, tag / use, notes)])
P = [
    ('P1', 'SCADA cabinet, 11 kV switchroom (controller pair, I/O rows 1 and 2, core network)', '11 kV switchroom', [
        ('CP8050', 2, 'CP-8050 A, CP-8050 B', 'Hot standby pair: A active, B standby.'),
        ('PS8642', 4, 'Two per controller', 'Confirm the CP-8050 accepts a second supply.'),
        ('CI8520', 2, 'One per controller', 'Separate port per network zone; not a switch.'),
        ('FLASH', 2, 'One per controller', ''),
        ('LIC', 2, 'One set per controller', 'Including the hot standby function if licensed separately (open point 12).'),
        ('CI8532', 1, 'Row 1 interface', 'Row 1 on its own interface so both controllers reach it.'),
        ('DI8112', 5, 'R1-DI01 to R1-DI05', '11 kV panels AH1 to AH8.'),
        ('DO8221', 2, 'R1-CO01, R1-CO02', 'Breaker open and close AH1 to AH7.'),
        ('DO8212', 1, 'R1-DO01', 'Alarm beacon, controller healthy.'),
        ('CI8532', 1, 'Row 2 interface', ''),
        ('DI8112', 2, 'R2-DI01, R2-DI02', 'Auxiliary TX, LV board, UPS, DC, E-stop, doors.'),
        ('AI8320', 1, 'R2-AI01', 'DC bus voltages.'),
        ('CORE', 2, 'SW-CORE-A, SW-CORE-B', ''),
        ('SFP', 4, 'Two per core switch', 'Ring ports and the core A-B link.'),
        ('FW', 1, 'FW-01', 'Second unit optional.'),
        ('GPS', 1, 'GPS-01', 'Antenna on the switchroom roof; cable length to suit.'),
        ('PSU20', 1, 'Cabinet 24 V', 'Switches, firewall, gateway supply. Runs from the 110 V DC station battery.'),
        ('PM5560', 1, 'PM5560-AUX', 'Mounted in the auxiliary LV board.'),
        ('FPP', 1, 'Fibre termination', ''),
        ('VX25', 1, 'SCADA cabinet', 'Door contact wired to row 2.'),
        ('ACC_C', 1, 'SCADA cabinet', ''),
    ]),
    ('P2', 'Control room (SCADA servers, operator and engineering stations)', 'Control room', [
        ('IPC', 2, 'SRV-A, SRV-B', 'Server A active, B hot standby.'),
        ('RACK', 1, 'Server rack', 'Not in the Device List Rev 1; needed to house the rack servers and UPS. Confirm the room.'),
        ('OWS', 1, 'OWS-01', ''),
        ('MON', 2, 'OWS-01', 'Two monitors.'),
        ('LAP', 1, 'ENG-01', ''),
        ('UPS', 1, 'UPS-01', 'SNMP to the SCADA.'),
        ('UPSB', 1, 'UPS-01', 'Size for the runtime the client requires.'),
        ('BEACON', 1, 'ALM-01', 'Hardwired to row 1.'),
    ]),
    ('P3', 'President station I/O panel (row 3, SW-PRS)', 'President office transformer station', [
        ('CI8533', 1, 'Row 3 interface', 'Confirm fibre type and maximum length.'),
        ('DI8112', 3, 'R3-DI01 to R3-DI03', 'RMU, 630 kVA transformer, LVDB.'),
        ('DO8221', 1, 'R3-CO01', 'RMU and LVDB commands.'),
        ('FSW', 1, 'SW-PRS', ''),
        ('SFP', 2, 'SW-PRS ring ports', ''),
        ('MGATE', 1, 'MGate-PRS', 'ST260E RMU relay, RS485.'),
        ('PSU5', 1, 'Panel 24 V', 'Smaller than the 20 A unit in the Device List, which is oversized here.'),
        ('PM5560', 1, 'PM5560-PRS', 'Mounted in the LVDB.'),
        ('FSB', 1, 'Fibre termination', ''),
        ('AX', 1, 'Row 3 panel', ''),
        ('ACC_P', 1, 'Row 3 panel', ''),
    ]),
    ('P4', 'Prime Minister station I/O panel (row 4, SW-PMS)', 'Prime Minister office transformer station', [
        ('CI8533', 1, 'Row 4 interface', 'Confirm fibre type and maximum length.'),
        ('DI8112', 3, 'R4-DI01 to R4-DI03', 'RMU, 630 kVA transformer, LVDB.'),
        ('DO8221', 1, 'R4-CO01', 'RMU and LVDB commands.'),
        ('FSW', 1, 'SW-PMS', ''),
        ('SFP', 2, 'SW-PMS ring ports', ''),
        ('MGATE', 1, 'MGate-PMS', 'ST260E RMU relay, RS485.'),
        ('PSU5', 1, 'Panel 24 V', ''),
        ('PM5560', 1, 'PM5560-PMS', 'Mounted in the LVDB.'),
        ('FSB', 1, 'Fibre termination', ''),
        ('AX', 1, 'Row 4 panel', ''),
        ('ACC_P', 1, 'Row 4 panel', ''),
    ]),
    ('P5', 'BESS and inverter transformer yard I/O panel (row 5, SW-BESS)', 'BESS yard', [
        ('CI8533', 1, 'Row 5 interface', ''),
        ('DI8112', 2, 'R5-DI01, R5-DI02', 'Fire panel, BESS ACB, PV ACB, MCCBs, inverter TX. Spare 19 %: see I/O Check.'),
        ('FSW', 1, 'SW-BESS', 'Mounted in the row 5 panel; Ethernet to the EMS in the container.'),
        ('SFP', 2, 'SW-BESS ring ports', ''),
        ('PSU5', 1, 'Panel 24 V', ''),
        ('FSB', 1, 'Fibre termination', ''),
        ('AX', 1, 'Row 5 panel', ''),
        ('ACC_P', 1, 'Row 5 panel', ''),
    ]),
    ('P6', 'Generator area I/O panel (row 6, SW-GEN)', 'Generator area', [
        ('CI8533', 1, 'Row 6 interface', ''),
        ('DI8112', 2, 'R6-DI01, R6-DI02', 'DG1, DG2, step-up TX, LV ACB.'),
        ('AI8320', 1, 'R6-AI01', 'Fuel level DG1, DG2.'),
        ('FSW', 1, 'SW-GEN', ''),
        ('SFP', 2, 'SW-GEN ring ports', ''),
        ('PSU5', 1, 'Panel 24 V', ''),
        ('MGATE', 0, 'MGate-GEN (option)', 'Only if the generator controllers are RS485 only (open point 3). Quantity 0 until confirmed.'),
        ('FSB', 1, 'Fibre termination', ''),
        ('AX', 1, 'Row 6 panel', ''),
        ('ACC_P', 1, 'Row 6 panel', ''),
    ]),
    ('P7', 'Inverter MV station network enclosure (SW-MV)', 'Inverter MV station', [
        ('FSW', 1, 'SW-MV', "PLC box and blue'Log XM."),
        ('SFP', 2, 'SW-MV ring ports', ''),
        ('PSU5', 1, 'Enclosure 24 V', ''),
        ('FSB', 1, 'Fibre termination', ''),
        ('AXS', 1, 'SW-MV enclosure', 'Not in the Device List Rev 1. Omit if the switch fits in the PLC box cabinet.'),
        ('ACC_P', 1, 'SW-MV enclosure', ''),
    ]),
    ('C', 'Cables and passive network material (site-wide; lengths from the cable route drawings)', 'Site', [
        ('FIBRE', None, '3 routes', 'Switchroom-President, switchroom-PM, President-PM; yard and generator area on spare cores. Length per route TBC.'),
        ('FCORD', 20, 'All fibre panels', ''),
        ('CAT6', None, 'Control room, switchroom, stations', 'Length TBC. Control room to switchroom assumed under 90 m; otherwise fibre.'),
        ('PCORD', 40, 'All panels', ''),
        ('RJ45', 30, 'Field connections', ''),
        ('RS485', None, 'MGate to ST260E', 'Short runs inside the stations. Length TBC.'),
        ('CTRL', None, 'Hardwired I/O', '244 hardwired signals. Length per cable schedule TBC.'),
    ]),
    ('S', 'Software and licences', 'SCADA servers and stations', [
        ('SCC', 1, 'SRV-A, SRV-B', '707 signals today plus about 30 % growth.'),
        ('SCCC', 1, 'OWS-01', 'Confirm whether the operator station runs as a client or as a web client.'),
        ('DM', 1, 'ENG-01', ''),
        ('WINS', 2, 'SRV-A, SRV-B', 'Usually supplied with the IPC.'),
        ('NMS', 1, 'SRV-A', 'Covers 7 switches, firewall, gateways, meters, UPS.'),
        ('FWLIC', 1, 'FW-01', 'VPN client (FortiClient) is free.'),
    ]),
    ('SP', 'Recommended spare parts (commissioning and first two years)', 'Store', [
        ('DI8112', 1, 'Spare', ''),
        ('DO8221', 1, 'Spare', ''),
        ('CI8533', 1, 'Spare', ''),
        ('PS8642', 1, 'Spare', ''),
        ('FLASH', 1, 'Spare', 'Keep a copy of the controller configuration on it.'),
        ('FSW', 1, 'Spare', ''),
        ('SFP', 2, 'Spare', ''),
        ('PSU5', 1, 'Spare', ''),
        ('MGATE', 1, 'Spare', ''),
    ]),
]

BY_OTHERS = [
    ('PV string inverters NB01 to NB04', 'Sungrow SG350HX', 'PV supplier', 'Power line communication to the PLC box; nothing connects to the SCADA directly.'),
    ('Power line communication box (PLC box)', 'Sungrow, model TBC', 'PV supplier', 'Ethernet port and Modbus TCP register map; written list of settings the SCADA may write.'),
    ('Embedded energy management system (EMS)', 'TBC', 'BESS / EMS supplier', 'Ethernet port, Modbus TCP point table, heartbeat, written list of settings the SCADA may write.'),
    ('Battery system 3,340 kWh', 'Sungrow ST3340UX', 'BESS supplier', 'Read through the EMS. Fire, E-stop and door volt-free contacts at a terminal strip.'),
    ('Generator controllers DG1, DG2', 'TBC', 'Generator supplier', 'Modbus TCP (or RTU, then add MGate-GEN); volt-free contacts for GCB, running, alarm, E-stop; 4-20 mA fuel level.'),
    ('11 kV protection relays AH1 to AH7', 'TBC', 'Switchgear supplier', 'Ethernet, two ports preferred; IEC 61850 or Modbus TCP documentation. Motorised breakers and auxiliary contacts.'),
    ('Neutral earthing resistor panel AH8', 'TBC', 'Switchgear supplier', 'Auxiliary contacts: breaker, earth switch, gas, earth fault.'),
    ('RMU protection relays', 'ST260E', 'RMU supplier', 'RS485 port and Modbus map; motorised switches.'),
    ('Transformers (inverter, DG step-up, auxiliary, 2 x 630 kVA)', 'TBC', 'Transformer supplier', 'Alarm and trip contacts wired to a terminal box.'),
    ('LV distribution boards and main breakers', 'TBC', 'Switchboard supplier', 'Space and CT for the PM5560 meters; breaker auxiliary contacts; motorised LVDB incomers.'),
    ('Fire detection and suppression panel', 'TBC', 'Fire system supplier', 'Volt-free contacts per the I/O list; direct trip of the battery outside the SCADA.'),
    ('UPS and 110 V DC station supplies', 'TBC', 'Electrical contractor', '110 V DC for the controllers and I/O rows; alarm contacts; Modbus TCP or SNMP.'),
    ("Weather station: blue'Log XM, SMP12-A, WindSonic, temperature sensors", 'Meteocontrol and others', 'PV supplier', 'Ethernet port and Modbus TCP documentation.'),
    ('RMU UPS 2 kVA', 'TBC', 'RMU supplier', 'On battery and fault contacts.'),
]

SERVICES = [
    ('Engineering: controller configuration, interlocks, IP plan, SCADA database and screens', 'lot'),
    ('Panel build: SCADA cabinet and six field panels, with wiring and internal tests', 'lot'),
    ('Factory acceptance test with the client', 'lot'),
    ('Site installation supervision, termination and fibre splicing', 'lot'),
    ('Site acceptance and interface tests with the PV, BESS, generator and switchgear suppliers', 'lot'),
    ('Training: operators and administrators', 'lot'),
    ('Documentation: as-built drawings, manuals, configuration backups', 'lot'),
    ('Warranty and remote support, 2 years', 'lot'),
]

IO_ROWS = [  # row, location, interface, modules (type, qty, channels), used signals (I/O list Rev 1)
    (1, '11 kV switchroom', 'CI-8532', [('DI-8112', 5, 16), ('DO-8221', 2, 12), ('DO-8212', 1, 8)], 87),
    (2, 'SCADA cabinet', 'CI-8532', [('DI-8112', 2, 16), ('AI-8320', 1, 4)], 25),
    (3, 'President station', 'CI-8533', [('DI-8112', 3, 16), ('DO-8221', 1, 12)], 43),
    (4, 'PM station', 'CI-8533', [('DI-8112', 3, 16), ('DO-8221', 1, 12)], 43),
    (5, 'BESS / inverter yard', 'CI-8533', [('DI-8112', 2, 16)], 26),
    (6, 'Generator area', 'CI-8533', [('DI-8112', 2, 16), ('AI-8320', 1, 4)], 20),
]


def head(ws, row, values, widths=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row, c, v)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, Alignment(wrap_text=True, vertical='center'), BOX
    ws.row_dimensions[row].height = 30
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def put(ws, row, values, fill=None, bold=False):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row, c, v)
        cell.alignment, cell.border = WRAP, BOX
        if fill:
            cell.fill = fill
        if bold:
            cell.font = Font(bold=True)


def key_of(k):
    return k


wb = Workbook()

# ---------------------------------------------------------------- Read Me
rm = wb.active
rm.title = 'Read Me'
rm.column_dimensions['A'].width = 24
rm.column_dimensions['B'].width = 110
lines = [
    ('Villa Somalia Hybrid Plant: SCADA Bill of Material', None),
    ('Project', 'Solarization of Villa Somalia, Mogadishu. 1,516 kWp PV, 3,340 kWh BESS, 2 diesel generators, 11 kV switchboard, 2 office transformer stations.'),
    ('Basis', 'SICAM A8000 design: SCADA Design Report Rev 1, Device and I/O List Rev 1, drawing set BFE/PR/26/01-SCADA-001 to 005.'),
    ('Status', 'Rev A, 22 Oct 2026, for pricing and supplier quotations. Quantities fixed by the design; prices to be entered.'),
    ('Architecture', 'Two CP-8050 in hot standby; six I/O rows (24 modules, 244 hardwired signals, 27 % spare); SICAM SCC on two servers by IEC 60870-5-104; '
                     'managed fibre ring of 7 switches; one firewall with VPN off by default; GPS time.'),
    ('Sheets', None),
    ('BOM by Panel', 'Every item grouped by the panel or area it goes into, with tag, quantity and notes. Use it to build and check the panels.'),
    ('Summary', 'The same items totalled per order line. Enter unit prices here, once; the BOM by Panel picks them up.'),
    ('IO Check', 'I/O modules per row against the hardwired signals of the I/O list, with spare capacity.'),
    ('By Others', 'Equipment that the SCADA connects to but does not supply, and what the SCADA needs from each supplier.'),
    ('Services', 'Engineering and site services, for the offer. Not materials.'),
    ('Number status key', None),
    (CHK, 'Ordering number found on a supplier listing or a Siemens configuration document.'),
    (MEM, 'Ordering code from knowledge of the maker\'s catalogue; not checked online. Confirm with the distributor.'),
    (SUP, 'Product exists but its number was not found, or it is configured or licensed by quotation. Take it from the supplier.'),
    (TBC, 'Size, make or length not fixed yet. Shaded yellow.'),
    ('Assumptions and points to confirm', None),
    ('1', 'Control voltage 110 V DC for the controllers and I/O rows (selects PS-8642, CI-8532, CI-8533, DI-8112). Confirm that 110 V DC exists at the office '
          'stations, the BESS yard and the generator area. If only 230 V AC is available there, either add a 230 V AC to 110 V DC supply per panel or change those '
          'rows to the 24 V variants (CI-8531, DI-8110) fed from the SITOP.'),
    ('2', 'Siemens to confirm: CI-8533 and DO-8221 ordering numbers, the hot standby function and licence, switchover time, and how the I/O rows connect to both controllers.'),
    ('3', 'Network ordering codes (Moxa, Fortinet, Meinberg) are from memory. Siemens SCALANCE equivalents may be offered instead; keep the same port counts.'),
    ('4', 'Field switches at the stations and yard use a 5 A SITOP instead of the 20 A unit of the Device List, which is oversized for one switch and one gateway.'),
    ('5', 'Two items not in the Device List Rev 1 are added: a 19-inch server rack with console for the control room, and a small enclosure for SW-MV.'),
    ('6', 'Cable and fibre lengths are left open until the cable route drawings are available.'),
    ('7', 'Currency and prices: enter in the Summary sheet. Totals calculate automatically.'),
]
for i, (a, b) in enumerate(lines, 1):
    rm.cell(i, 1, a)
    if b:
        rm.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical='top')
    if b is None:
        rm.cell(i, 1).font = Font(bold=True, size=14 if i == 1 else 11, color=NAVY)
    if a == TBC:
        rm.cell(i, 1).fill = TBC_FILL

# ---------------------------------------------------------------- Summary (built first so its rows are known)
sm = wb.create_sheet('Summary')
bp = wb.create_sheet('BOM by Panel')
wb.move_sheet('BOM by Panel', -1)  # BOM by Panel before Summary

order, used = [], {}
for code, _, _, items in P:
    for k, *_ in items:
        if k not in order:
            order.append(k)
        if code not in used.setdefault(k, []):
            used[k].append(code)
SM0 = 3
sm['A1'] = 'Summary by order line'
sm['A1'].font = Font(bold=True, size=14, color=NAVY)
sm['H1'] = 'Currency:'
sm['I1'] = '[USD]'
sm['I1'].fill = TBC_FILL
head(sm, 2, ['Line', 'Description', 'Manufacturer', 'Model', 'Order no.', 'Number status', 'Unit', 'Total qty', 'Unit price', 'Total price', 'Used in', 'Key'],
     [6, 60, 18, 24, 22, 22, 7, 10, 13, 15, 20, 4])
smrow = {}
n_bp = sum(len(it) + 3 for *_, it in P) + 10
for i, k in enumerate(order):
    r = SM0 + i
    d = CAT[k]
    smrow[k] = r
    put(sm, r, [i + 1, d[0], d[1], d[2], d[3], d[4], d[5],
                "=IF(COUNTIFS('BOM by Panel'!$O$3:$O$%d,L%d,'BOM by Panel'!$H$3:$H$%d,\"<>\")=0,\"TBC\",SUMIF('BOM by Panel'!$O$3:$O$%d,L%d,'BOM by Panel'!$H$3:$H$%d))"
                % (n_bp, r, n_bp, n_bp, r, n_bp),
                None, '=IF(OR(I%d="",H%d="TBC"),"",H%d*I%d)' % (r, r, r, r), ', '.join(used[k]), key_of(k)])
    sm.cell(r, 9).fill = TBC_FILL
    if d[4] in (TBC,):
        sm.cell(r, 6).fill = TBC_FILL
    for c in (9, 10):
        sm.cell(r, c).number_format = '#,##0.00'
last = SM0 + len(order) - 1
put(sm, last + 1, ['', 'Total materials', '', '', '', '', '', '', '', '=SUM(J%d:J%d)' % (SM0, last), '', ''], bold=True)
sm.cell(last + 1, 10).number_format = '#,##0.00'
sm.column_dimensions['L'].hidden = True
sm.freeze_panes = 'C3'

# ---------------------------------------------------------------- BOM by Panel
bp['A1'] = 'Bill of material by panel – SICAM A8000 design'
bp['A1'].font = Font(bold=True, size=14, color=NAVY)
head(bp, 2, ['Item', 'Tag / use', 'Description', 'Manufacturer', 'Model', 'Order no.', 'Number status', 'Qty', 'Unit', 'Unit price', 'Total price',
             'Location', 'Notes', 'Group', 'Key'],
     [8, 22, 48, 16, 22, 20, 20, 7, 6, 12, 14, 20, 52, 7, 4])
r = 3
for code, title, loc, items in P:
    put(bp, r, [code, title] + [''] * 13, fill=SEC_FILL, bold=True)
    bp.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
    first = r + 1
    r += 1
    for j, (k, q, tag, note) in enumerate(items, 1):
        d = CAT[k]
        sr = smrow[k]
        put(bp, r, ['%s.%02d' % (code, j), tag, d[0], d[1], d[2], d[3], d[4], q, d[5],
                    "=IF(Summary!$I$%d=\"\",\"\",Summary!$I$%d)" % (sr, sr), '=IF(OR(J%d="",H%d=""),"",H%d*J%d)' % (r, r, r, r),
                    loc, note, code, key_of(k)])
        if q is None:
            bp.cell(r, 8).fill = TBC_FILL
        if d[4] == TBC:
            bp.cell(r, 7).fill = TBC_FILL
        for c in (10, 11):
            bp.cell(r, c).number_format = '#,##0.00'
        r += 1
    put(bp, r, ['', 'Subtotal ' + code] + [''] * 8 + ['=SUM(K%d:K%d)' % (first, r - 1)] + [''] * 4, bold=True)
    bp.cell(r, 11).number_format = '#,##0.00'
    r += 2
put(bp, r, ['', 'Total materials'] + [''] * 8 + ['=Summary!J%d' % (last + 1)] + [''] * 4, bold=True)
bp.cell(r, 11).number_format = '#,##0.00'
assert r < n_bp, (r, n_bp)
bp.column_dimensions['O'].hidden = True
bp.freeze_panes = 'C3'

# ---------------------------------------------------------------- I/O Check
io = wb.create_sheet('IO Check')
io['A1'] = 'I/O modules against the hardwired signals of the I/O list Rev 1'
io['A1'].font = Font(bold=True, size=14, color=NAVY)
head(io, 2, ['Row', 'Location', 'Interface', 'Modules', 'Module count', 'Channels', 'Signals used', 'Spare', 'Spare %', 'Remark'],
     [6, 24, 12, 44, 13, 10, 13, 9, 9, 60])
for i, (row, loc, ci, mods, used) in enumerate(IO_ROWS):
    rr = 3 + i
    put(io, rr, [row, loc, ci, ', '.join('%s x%d' % (m, q) for m, q, _ in mods), sum(q for _, q, _ in mods),
                 sum(q * ch for _, q, ch in mods), used, '=F%d-G%d' % (rr, rr), '=H%d/F%d' % (rr, rr),
                 'Below the 20 % target: add a third DI-8112 if more yard signals appear.' if row == 5 else ''])
    io.cell(rr, 9).number_format = '0%'
t = 3 + len(IO_ROWS)
put(io, t, ['', 'Total', '', '', '=SUM(E3:E%d)' % (t - 1), '=SUM(F3:F%d)' % (t - 1), '=SUM(G3:G%d)' % (t - 1), '=SUM(H3:H%d)' % (t - 1),
            '=H%d/F%d' % (t, t), ''], bold=True)
io.cell(t, 9).number_format = '0%'
put(io, t + 2, ['', 'Check', '', 'I/O modules in the BOM (rows 1 to 6, without spares)',
                "=SUMPRODUCT(('BOM by Panel'!$E$3:$E$%d={\"DI-8112\",\"DO-8221\",\"DO-8212\",\"AI-8320\"})*('BOM by Panel'!$N$3:$N$%d<>\"SP\")*('BOM by Panel'!$H$3:$H$%d))"
                % (n_bp, n_bp, n_bp), '', '', '', '', '=IF(E%d=E%d,"Matches the I/O list","Check the module quantities")' % (t + 2, t)])

# ---------------------------------------------------------------- By Others
bo = wb.create_sheet('By Others')
bo['A1'] = 'Equipment the SCADA connects to but does not supply'
bo['A1'].font = Font(bold=True, size=14, color=NAVY)
head(bo, 2, ['No.', 'Equipment', 'Make / model', 'Supplied by', 'What the SCADA needs from the supplier'], [6, 48, 26, 24, 90])
for i, row in enumerate(BY_OTHERS):
    put(bo, 3 + i, [i + 1] + list(row))
    if 'TBC' in row[1]:
        bo.cell(3 + i, 3).fill = TBC_FILL

# ---------------------------------------------------------------- Services
sv = wb.create_sheet('Services')
sv['A1'] = 'Engineering and services (for the offer; not materials)'
sv['A1'].font = Font(bold=True, size=14, color=NAVY)
head(sv, 2, ['No.', 'Service', 'Unit', 'Qty', 'Price'], [6, 90, 8, 8, 14])
for i, (s, u) in enumerate(SERVICES):
    put(sv, 3 + i, [i + 1, s, u, 1, None])
    sv.cell(3 + i, 5).fill = TBC_FILL
    sv.cell(3 + i, 5).number_format = '#,##0.00'
e = 3 + len(SERVICES)
put(sv, e, ['', 'Total services', '', '', '=SUM(E3:E%d)' % (e - 1)], bold=True)
sv.cell(e, 5).number_format = '#,##0.00'

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print('wrote', OUT, '·', len(order), 'order lines')
