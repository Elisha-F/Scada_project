"""Rev 2 of the Villa Somalia SCADA Design Report: control platform changed to Schneider Electric.

Input is the Rev 1 report (Siemens SICAM A8000, from tools/revise_report.py). Each edit names the exact
old text, so a changed source fails loudly instead of editing the wrong place. A final sweep fails if
any Siemens-platform term is left.

Usage: python3 tools/convert_report_schneider.py <rev1.docx> <rev2.docx>
"""
import os
import re
import sys
from copy import deepcopy
import docx
from PIL import Image

SRC, OUT = sys.argv[1], sys.argv[2]
HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, '..', 'docs', 'figures')
COMMS = 462
d = docx.Document(SRC)
T = d.tables


def all_pars():
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def set_par(p, old, new):
    """Replace old with new inside one paragraph, keeping the run formatting around it."""
    text = ''.join(r.text for r in p.runs)
    assert old in text, (old, text)
    start = text.index(old)
    pos, first = 0, None
    for r in p.runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= start or a >= start + len(old):
            continue
        if first is None:
            first = r
            r.text = r.text[:start - a] + new + (r.text[start + len(old) - a:] if b >= start + len(old) else '')
        else:
            r.text = r.text[start + len(old) - a:] if b > start + len(old) else ''


def par(old, new):
    ps = [p for p in d.paragraphs if old in p.text]
    assert len(ps) == 1, (old, len(ps))
    set_par(ps[0], old, new)
    return ps[0]


def cell(t, r, c, old, new):
    ps = [p for p in T[t].rows[r].cells[c].paragraphs if old in p.text]
    assert ps, ('T%d R%d C%d' % (t, r, c), old, T[t].rows[r].cells[c].text)
    set_par(ps[0], old, new)


def gsub(old, new, at_least=1):
    n = 0
    for p in all_pars():
        while old in ''.join(r.text for r in p.runs):
            set_par(p, old, new)
            n += 1
    assert n >= at_least, (old, n)
    return n


def set_cell_text(c, v):
    p = c.paragraphs[0]
    for extra in c.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if not p.runs:
        p.add_run('')
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = str(v)


def rebuild(ti, rows, header=None):
    """Keep the header row, use row 1 as the template and replace every body row."""
    t = T[ti]
    if header:
        for c, v in zip(t.rows[0].cells, header):
            set_cell_text(c, v)
    tmpl = deepcopy(t.rows[1]._tr)
    for row in list(t.rows)[1:]:
        t._tbl.remove(row._tr)
    for vals in rows:
        tr = deepcopy(tmpl)
        t._tbl.append(tr)
        for c, v in zip(t.rows[-1].cells, vals):
            set_cell_text(c, v)


def add_row(ti, vals):
    tr = deepcopy(T[ti].rows[-1]._tr)
    T[ti]._tbl.append(tr)
    for c, v in zip(T[ti].rows[-1].cells, vals):
        set_cell_text(c, v)


def insert_par_after(p, text):
    el = deepcopy(p._p)
    p._p.addnext(el)
    np = docx.text.paragraph.Paragraph(el, p._parent)
    for r in np.runs[1:]:
        r._r.getparent().remove(r._r)
    np.runs[0].text = text
    return np


# ---------------------------------------------------------------- 1 Introduction and references
p = par('It answers the client\'s review comments on the earlier drawing set.', 'It answers the client\'s review comments on the earlier drawing set.')
insert_par_after(p, 'Revision 2 changes the control platform from Siemens SICAM A8000 to Schneider Electric: Modicon M580 hot standby controllers, '
                    'Modicon X80 and Edge I/O NTS input and output modules, and EcoStruxure Machine SCADA Expert. The plant, the signals and tags, '
                    'the EMS rules, the network and the cybersecurity design are unchanged except where the platform requires.')
cell(0, 3, 0, 'Rev1.xlsx', 'Rev2.xlsx')
cell(0, 3, 1, 'I/O rows', 'I/O locations')
add_row(0, ['Offer Proposal Villa Somalia, Schneider Electric, version 1, 16 Oct 2026', 'Vendor offer for the controllers, I/O and SCADA; to be requoted to this design'])

# ---------------------------------------------------------------- 3 Design basis
par('This sets the I/O module variants.', 'Field contacts are wetted at 24 V DC from the I/O panels, and commands switch the 110 V DC circuits through interposing relays.')
par('The 11 kV relays offer IEC 61850 or Modbus TCP.', 'The 11 kV relays offer Modbus TCP. IEC 61850 would need an IEC 61850 module on the controllers.')
cell(4, 1, 0, 'IEC 61850', 'Modbus TCP and EtherNet/IP')
cell(4, 1, 1, 'Communication with the 11 kV protection relays', 'Links between the controllers, the SCADA stations, the field devices and the Edge I/O islands')
cell(4, 2, 0, 'IEC 60870-5-104', 'IEC 61131-3')
cell(4, 2, 1, 'Link between the controller and the SCADA servers', 'Controller programming in EcoStruxure Control Expert')

# ---------------------------------------------------------------- 4 Architecture
par('A hot standby pair of Siemens SICAM A8000 CP-8050 controllers collects every signal in the plant and passes it to two SCADA servers.',
    'A hot standby pair of Schneider Electric Modicon M580 controllers collects every signal in the plant and passes it to two SCADA stations.')
par('It reads five groups of devices over the network and six rows of hardwired signals.',
    'It reads five groups of devices over the network and hardwired signals at five I/O locations: a remote drop in the switchroom and four Edge I/O islands.')
par('5 network sources, 6 I/O rows', '5 network sources, 5 I/O locations')
par('and the command module checks that only one output operates.',
    'and the controller allows only one command output per device at a time; the output drives an interposing relay in the breaker control circuit.')
rebuild(5, [
    ['Modicon M580 hot standby pair A and B (BMEH582040)', 'Plant controller: collects all signals, time-stamps events, executes breaker commands, passes settings to the EMS; B takes over if A fails',
     '11 kV switchroom, in the SCADA cabinet beside the switchboard'],
    ['SCADA stations A and B (EcoStruxure Machine SCADA Expert)', 'Screens, alarms, trends, reports, user logins and audit log; also the operator stations; B takes over if A fails', 'Control room'],
    ['Engineering laptop', 'Controller and SCADA configuration and diagnostics (EcoStruxure Control Expert)', 'Control room'],
    ['Embedded EMS', 'Dispatch of PV, battery and generators', 'BESS area'],
    ['Protection relays', 'Protection of the 11 kV and RMU circuits', '11 kV switchboard and office stations'],
    ['Switchroom remote drop and four Edge I/O islands', 'Hardwired signals from the field', 'Near the equipment they serve'],
])

# ---------------------------------------------------------------- 5 Scope, 6 PV and BESS interface
par('Relays report over IEC 61850 or Modbus TCP; positions and commands are also hardwired to row 1.',
    'Relays report over Modbus TCP; positions and commands are also hardwired to the switchroom I/O.')
par('breaker positions are hardwired to rows 3 and 4.', 'breaker positions are hardwired to the office station I/O islands.')
par('critical states are duplicated by contacts to row 6.', 'critical states are duplicated by contacts to the generator area I/O island.')
par('safety states are hardwired to row 5.', 'safety states are hardwired to the BESS yard I/O island.')
par('through the relays and I/O row 1.', 'through the relays and the switchroom I/O.')
cell(7, 5, 1, 'IEC 61850 or Modbus TCP from the relays', 'Modbus TCP from the relays')
cell(7, 6, 1, 'Modbus RTU from the ST260E relays, through gateways', 'Modbus RTU from the ST260E relays, through the Edge I/O serial masters')
cell(8, 1, 2, 'PLC box to CP-8050', 'PLC box to the controllers')
cell(8, 2, 2, 'EMS and CP-8050, both ways', 'EMS and the controllers, both ways')
cell(8, 3, 2, 'to I/O row 5', 'to the BESS yard I/O island')
cell(9, 7, 1, 'CP-8050 and secured command outputs', 'Controllers and interposing relays')
cell(10, 4, 0, 'Hardwired, I/O row 5', 'Hardwired, BESS yard I/O island')
cell(11, 6, 2, 'Wires them to I/O row 5', 'Wires them to the BESS yard I/O island')

# ---------------------------------------------------------------- 7 I/O summary
par('The plant has 244 hardwired signals on 24 I/O modules in six rows, with 27% spare channels. Row 5, in the BESS and inverter yard, is the fullest at 19% spare;',
    'The plant has 244 hardwired signals on 19 I/O modules at five locations, with 32% spare channels (28% on digital inputs). The BESS and inverter yard is the fullest at 19% spare;')
par('A further 463 signals arrive', 'A further %d signals arrive' % COMMS)
ps = [p for p in d.paragraphs if p.text == 'I/O rows']
assert len(ps) == 1
set_par(ps[0], 'I/O rows', 'I/O locations')
par('Each row is a panel of I/O modules placed near the equipment it serves. Field contacts are wired a short distance to that panel, and one link carries all its signals to the controllers. Both controllers reach every row; only the active one reads inputs and drives outputs.',
    'Each location has its own I/O near the equipment it serves: a remote drop of Modicon X80 modules in the switchroom, and Modicon Edge I/O NTS islands at the two office stations, the BESS yard and the generator area. '
    'Field contacts are wired a short distance to the local I/O. The switchroom drop is on the M580 Ethernet remote I/O network; the islands connect over the fibre ring by EtherNet/IP. '
    'Both controllers reach every location; only the active one reads inputs and drives outputs.')
par('The controller accepts up to 16 rows of 8 modules, so 10 rows remain free for expansion.',
    'Further remote drops and Edge I/O islands can be added for new locations without changing the controllers, and each location keeps spare channels for new signals.')
rebuild(12, [
    ['SWR', '11 kV switchroom, SCADA cabinet', 'X80 remote drop', 'Copper, remote I/O', '11 kV panels, beacon, auxiliaries, UPS, DC, E-stop, doors', '6', '112'],
    ['PRS', 'President office station', 'Edge I/O island', 'Fibre ring, EtherNet/IP', 'RMU, transformer, LVDB, RMU UPS', '4', '43'],
    ['PMS', 'Prime Minister office station', 'Edge I/O island', 'Fibre ring, EtherNet/IP', 'RMU, transformer, LVDB, RMU UPS', '4', '43'],
    ['YRD', 'BESS and inverter transformer yard', 'Edge I/O island', 'Fibre ring, EtherNet/IP', 'Inverter transformer, LV breakers, fire panel, E-stop, door', '2', '26'],
    ['GEN', 'Generator area', 'Edge I/O island', 'Fibre ring, EtherNet/IP', 'Generators, fuel, step-up transformer', '3', '20'],
    ['Total', '', '', '', '', '19', '244'],
], header=['Location', 'Where', 'I/O hardware', 'Link', 'Equipment wired', 'Modules', 'Signals'])
rebuild(13, [
    ['BMXDDI3202K', 'X80 digital input, 32 channels, 24 V DC (switchroom)', '4', '128', '94', '34'],
    ['BMXDDO3202K', 'X80 digital output, 32 channels, 24 V DC (switchroom)', '1', '32', '16', '16'],
    ['BMXAMI0410', 'X80 analogue input, 4 isolated channels (switchroom)', '1', '4', '2', '2'],
    ['NTSDDI1602XK', 'Edge I/O digital input, 16 channels, 24 V DC', '10', '160', '114', '46'],
    ['NTSDDO1602K', 'Edge I/O digital output, 16 channels, 0.5 A', '2', '32', '16', '16'],
    ['To select', 'Edge I/O analogue input, 4 channels (generator area)', '1', '4', '2', '2'],
    ['Total', '', '19', '360', '244', '116'],
])
par('Breaker commands use the secured command module, which checks that only one output operates and that the coil circuit is intact before switching.',
    'Breaker commands drive interposing relays whose contacts switch the 110 V DC control circuit; the 24 V DC outputs never switch a coil directly. '
    'The controller allows only one command per device at a time, and the protection relays supervise the trip circuits.')

# ---------------------------------------------------------------- 8 Communication
par('The SCADA servers, workstation, firewall, both controllers and the 11 kV relays each connect to both core switches.',
    'The SCADA stations, engineering laptop, firewall, both controllers and the 11 kV relays each connect to both core switches.')
cell(15, 1, 0, 'Controllers A and B to SCADA servers A and B', 'Controllers A and B to SCADA stations A and B')
cell(15, 1, 1, 'IEC 60870-5-104', 'Modbus TCP (EtherNet/IP as an alternative)')
cell(15, 2, 1, 'IEC 61850, or Modbus TCP', 'Modbus TCP; IEC 61850 with an added controller module')
cell(15, 7, 1, 'Modbus RTU through a station gateway', 'Modbus RTU through the Edge I/O serial master')
cell(15, 9, 0, 'Controllers to I/O rows 1 and 2', 'Controllers to the switchroom I/O drop')
cell(15, 9, 1, 'SICAM I/O link', 'M580 Ethernet remote I/O')
cell(15, 9, 3, 'None; same cabinet', 'Both controllers on the remote I/O network')
cell(15, 10, 0, 'Controller to I/O rows 3 to 6', 'Controllers to the Edge I/O islands')
cell(15, 10, 1, 'SICAM I/O link, optical', 'EtherNet/IP')
cell(15, 10, 2, 'Dedicated fibre pair', 'Fibre ring')
cell(15, 10, 3, 'None; see below', 'Ring')
cell(15, 14, 1, 'Redundancy synchronisation (Siemens)', 'M580 hot standby link')
cell(15, 14, 2, 'Direct cable between the controllers, to confirm', 'Dedicated link between the CPUs')
cell(15, 14, 3, 'Through the core switches, to confirm', 'None; Schneider to confirm the behaviour on loss of the link')
cell(16, 8, 0, 'Fibre pair to a remote I/O row is cut', 'The field or station switch at an I/O island fails')
cell(16, 8, 1, 'Hardwired signals from that location are lost until repaired; network data from that location still arrives through the ring',
     'Hardwired signals from that location are lost until repaired; the ring keeps every other location')
par('Remote I/O row links.', 'I/O islands.')
par('Each runs on one fibre pair. Siemens will confirm whether the optical row link can run as a ring; if not, a small CP-8031 controller at a site would put its signals on the ring.',
    'Each island connects through its local field or station switch; if that switch fails, that location\'s hardwired signals are lost until it is repaired, while the ring keeps every other location.')
par('Each I/O row interface remains single; its failure loses only that row.', 'Each I/O location has one network interface; its failure loses only that location.')

# ---------------------------------------------------------------- 9 Cybersecurity
par('by separate ports on the controller,', 'by separate Ethernet modules on the controller racks (BMENOC0321 forwarding only permitted traffic),')
par('by firewall rules that deny all traffic not listed in section 8.',
    'by firewall rules that deny all traffic not listed in section 8. The M580 controllers carry built-in security functions to IEC 62443 and Achilles Level 2 certification, according to the manufacturer.')
cell(17, 1, 1, 'SCADA servers, workstation, engineering laptop', 'SCADA stations, engineering laptop')
cell(17, 3, 1, 'Relays, power meters, gateways, I/O rows', 'Relays, power meters, switchroom I/O drop, Edge I/O islands')

# ---------------------------------------------------------------- 10 Switching rules, 11 Sequences, 12 HMI
par('The secured command module checks that only one output operates and that the coil circuit is intact.',
    'The controller allows only one command output per device at a time; the command drives an interposing relay, and the protection relay supervises the trip circuit.')
par('does not depend on the CP-8050, the SCADA or the network', 'does not depend on the controllers, the SCADA or the network')
par('The CP-8050 monitors the EMS heartbeat', 'The controllers monitor the EMS heartbeat')
par('the final screens are built in SICAM SCC', 'the final screens are built in EcoStruxure Machine SCADA Expert')
cell(23, 13, 1, 'I/O row and server', 'I/O location and SCADA station')
cell(23, 15, 1, 'I/O row, module and channel', 'I/O location, module and channel')

# ---------------------------------------------------------------- 13 Equipment
par('The SCADA uses Siemens for the controller, I/O and SCADA software, and industrial-grade network and power equipment. Datasheets for each item are collected in Attachment C. Product numbers marked "to confirm" will be taken from the supplier quotations.',
    'The SCADA uses Schneider Electric for the controllers, I/O, SCADA software, network switches, power meters, power supplies, UPS and enclosures; the firewall, GPS time server and engineering laptop are from other makers. '
    'Datasheets for each item are collected in Attachment C. References marked "to select" or "to confirm" will be taken from the Schneider quotation, which is to be requoted to this design.')
par('Controller and I/O (Siemens SICAM A8000)', 'Controller and I/O (Schneider Electric Modicon M580, X80 and Edge I/O)')
rebuild(24, [
    ['Hot standby CPU, level 2', 'Modicon M580', 'BMEH582040', '2'],
    ['Ethernet backplane, 6 slots, dual power supply (CPU racks)', 'Modicon X80', 'BMEXBP0602', '2'],
    ['Redundant power supply, 24-48 V DC (CPU racks)', 'Modicon X80', 'To select', '4'],
    ['Ethernet communication module, device network', 'Modicon M580', 'BMENOC0311', '2'],
    ['Ethernet communication module with IP forwarding, control network', 'Modicon M580', 'BMENOC0321', '2'],
    ['Hot standby link transceivers', 'Modicon M580', 'To select', '1 set'],
    ['Ethernet backplane, 12 slots (switchroom drop)', 'Modicon X80', 'BMEXBP1200', '1'],
    ['Power supply, 24-48 V DC (switchroom drop)', 'Modicon X80', 'BMXCPS3020', '1'],
    ['Ethernet remote drop adapter', 'Modicon X80', 'BMECRA31210', '1'],
    ['Digital input, 32 channels, with pre-wired cable', 'Modicon X80', 'BMXDDI3202K, BMXFCW303', '4'],
    ['Digital output, 32 channels, with pre-wired cable', 'Modicon X80', 'BMXDDO3202K, BMXFCW303', '1'],
    ['Analogue input, 4 isolated channels', 'Modicon X80', 'BMXAMI0410', '1'],
    ['Serial link module, Modbus RTU', 'Modicon X80', 'BMXNOM0200', '1'],
    ['Edge I/O network interface and power supply', 'Modicon Edge I/O NTS', 'NTSNEC1200K, NTSPFB1002HK', '4 each'],
    ['Edge I/O digital input, 16 channels', 'Modicon Edge I/O NTS', 'NTSDDI1602XK', '10'],
    ['Edge I/O digital output, 16 channels', 'Modicon Edge I/O NTS', 'NTSDDO1602K', '2'],
    ['Edge I/O analogue input, 4 channels', 'Modicon Edge I/O NTS', 'To select', '1'],
    ['Edge I/O serial master, Modbus RTU (ST260E)', 'Modicon Edge I/O NTS', 'NTSFMB0120K', '2'],
    ['Edge I/O terminal blocks', 'Modicon Edge I/O NTS', 'NTSXTB02030H, NTSXTB18001XH, NTSXTB18001H, NTSXTB12001H', '22'],
    ['Interposing relays, 24 V DC coil, contact for 110 V DC', 'Zelio Relay', 'To select', '32'],
    ['Controller engineering software', 'EcoStruxure Control Expert', 'By quotation', '1'],
])
rebuild(25, [
    ['SCADA stations, active and hot standby, industrial panel PC', 'Schneider Electric HMIP67GD0F88BE7N00 (model to confirm)', '2'],
    ['SCADA runtime licence, 4K tags', 'Schneider Electric EcoStruxure Machine SCADA Expert, HMIEMSE3PRT4KA', '2'],
    ['SCADA build-time licence, 4K tags', 'Schneider Electric EcoStruxure Machine SCADA Expert, HMIEMSEBT4KA', '1'],
    ['Engineering laptop', 'Dell Latitude 5430 Rugged', '1'],
    ['Core switches, managed, with SFP fibre ports', 'Schneider Electric Modicon managed switch (reference to select)', '2'],
    ['Field and station switches, managed, with SFP fibre ports', 'Schneider Electric Modicon managed switch (reference to select)', '5'],
    ['Fibre transceivers, single-mode', 'Schneider Electric SFP for the switches (to select)', '14'],
    ['Firewall with VPN', 'Fortinet FortiGate Rugged 60F', '1'],
    ['GPS time server', 'Meinberg LANTIME M300', '1'],
    ['Power meters for LVDBs and auxiliary board', 'Schneider Electric PowerLogic PM5560', '3'],
    ['24 V DC power supplies', 'Schneider Electric Phaseo (rating to select)', '6'],
    ['UPS for the SCADA cabinet', 'APC Smart-UPS SRT 3000 VA (Schneider Electric)', '1'],
    ['SCADA cabinet', 'Schneider Electric Spacial SF, 800 x 2000 x 800 mm', '1'],
    ['Outdoor I/O panels, IP55 minimum', 'Schneider Electric Thalassa PLM (size to suit)', '4'],
    ['Alarm sounder and beacon', 'Schneider Electric Harmony XVS (model to select)', '1'],
    ['Armoured single-mode fibre, 12 core', 'To suit', '3 routes'],
])
rebuild(26, [
    ['Memory card with the configured controller application', '1', 'Controller B is the running spare; the card restores a replaced CPU quickly'],
    ['BMXDDI3202K, BMXDDO3202K, BMXAMI0410, BMXCPS3020', '1 each', 'One of each X80 module type on site'],
    ['NTSNEC1200K, NTSPFB1002HK, NTSDDI1602XK, NTSDDO1602K, NTSFMB0120K', '1 each', 'One of each Edge I/O module type on site'],
    ['Interposing relays', '2', 'Command circuit repair'],
    ['Field switch and fibre transceivers', '1 and 2', 'Network repair'],
])
par('The controllers are supplied as a hot standby pair. I/O row 1 has its own CI-8532 interface rather than sitting on the controller backplane, so both controllers reach it. Siemens will confirm the redundancy function and licence, the switchover time, and how the I/O rows connect to both controllers.',
    'The controllers are an M580 hot standby pair: two BMEH582040 CPUs, each in its own rack with two power supplies, joined by the hot standby link. No I/O sits in the CPU racks; '
    'the switchroom I/O is in a remote drop and the other locations use Edge I/O islands, so both controllers reach every signal. Schneider will confirm the switchover time and the behaviour of the drop and islands on switchover.')
par('Siemens SICAM A8000: CP-8050, PS-8642, CI-8520, CI-8532, CI-8533, DI-8112, DO-8221, DO-8212, AI-8320',
    'Schneider Electric Modicon M580 and X80: BMEH582040, BMEXBP0602, BMEXBP1200, BMXCPS3020, BMENOC0311, BMENOC0321, BMECRA31210, BMXDDI3202K, BMXDDO3202K, BMXAMI0410, BMXNOM0200')
par('Siemens SIMATIC IPC647E and SICAM SCC',
    'Schneider Electric Modicon Edge I/O NTS: NTSNEC1200K, NTSPFB1002HK, NTSDDI1602XK, NTSDDO1602K, NTSFMB0120K and the analogue input module')
par('Moxa EDS-G512E, EDS-510E, SFP-1GLXLC, MGate MB3170',
    'Schneider Electric EcoStruxure Machine SCADA Expert, EcoStruxure Control Expert, industrial panel PC HMIP67GD0F88BE7N00, managed switches and SFP transceivers')
par('Schneider Electric PM5560; Siemens SITOP PSU8200; APC Smart-UPS SRT',
    'Schneider Electric PowerLogic PM5560, Phaseo power supplies, Zelio interface relays, Spacial and Thalassa enclosures, Harmony XVS beacon; APC Smart-UPS SRT')

# ---------------------------------------------------------------- 14 Open points and attachments
cell(27, 3, 3, 'Row 6 signals and analogue inputs', 'Generator area signals and analogue inputs')
cell(27, 4, 3, 'Protocol licences; possible gateway', 'IEC 61850 module if the relays are read by IEC 61850')
cell(27, 7, 3, 'Command outputs at rows 3 and 4', 'Command outputs at the office stations')
cell(27, 8, 3, 'Fibre or copper for remote rows; fibre lengths', 'Fibre lengths to the office stations, BESS yard and generator area')
cell(27, 10, 1, 'Product numbers for DO-8221, CI-8533 and the protocol licences; fibre limits for CI-8533',
     'References still to select: Edge I/O analogue input, interposing relays, managed switches with SFP ports, hot standby link transceivers, redundant power supplies, 24 V supplies')
cell(27, 10, 2, 'Siemens', 'Schneider Electric')
cell(27, 11, 3, 'Row 6', 'Generator area I/O')
cell(27, 12, 1, 'Hot standby controller pair: redundancy function and licence, switchover time, connection of the I/O rows to both controllers, product numbers',
     'M580 hot standby: switchover time, behaviour of the switchroom drop and Edge I/O islands on switchover, loss of the hot standby link')
cell(27, 12, 2, 'Siemens', 'Schneider Electric')
cell(27, 12, 3, 'Sections 4, 7, 8, 13', 'Sections 4, 7, 8')
cell(27, 16, 1, 'behaviour on loss of the CP-8050 link', 'behaviour on loss of the controller link')
add_row(27, ['19', 'SCADA redundancy between the two Machine SCADA Expert stations, historian location and retention', 'Schneider Electric', 'Sections 4, 10'])
add_row(27, ['20', 'Optional panel servers with humidity and temperature sensors (PAS600, CL110, TH110) offered by Schneider; not in the I/O list', 'Client decision', 'Section 13'])
add_row(27, ['21', 'Requote of the Schneider offer of 16 Oct 2026 to this design: hot standby, I/O per location, fibre, interposing relays; relays and meters to be clarified', 'Schneider Electric', 'Sections 4, 7, 13'])
cell(28, 5, 2, 'final screens built in SICAM SCC', 'final screens built in EcoStruxure Machine SCADA Expert')

# ---------------------------------------------------------------- everywhere
gsub('Rev1.xlsx', 'Rev2.xlsx')
gsub('463 communication points', '%d communication points' % COMMS)
cell(14, 4, 1, '463', str(COMMS))
gsub('SCADA servers', 'SCADA stations')
gsub('Server B takes over', 'Station B takes over')
gsub('Repair server A', 'Repair station A')
gsub('SCADA server A', 'SCADA station A')
gsub(' to row 1', ' to the switchroom I/O')
gsub(' to row 2', ' to the switchroom I/O')
gsub(' to rows 3 and 4', ' to the office station I/O islands')
gsub(' to row 5', ' to the BESS yard I/O island')
gsub(' to row 6', ' to the generator area I/O island')
gsub('to the nearest row', 'to the nearest I/O location')
gsub('SCADA server failure', 'SCADA station failure')
gsub('standby controller or server;', 'standby controller or SCADA station;')
gsub('Modbus TCP, or Modbus RTU through a gateway', 'Modbus TCP, or Modbus RTU through an Edge I/O serial master')
cell(18, 7, 0, 'Workstations', 'SCADA stations and laptop')

# ---------------------------------------------------------------- figures 1 to 3
for k, name in ((0, 'fig1_architecture_se.png'), (1, 'fig2_pv_bess_interface_se.png'), (2, 'fig3_network_se.png')):
    shape = d.inline_shapes[k]
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    path = os.path.join(FIG, name)
    d.part.related_parts[rid]._blob = open(path, 'rb').read()
    w, h = Image.open(path).size
    shape.height = int(shape.width * h / w)

# section 12 mock-ups from the converted HMI (tools/capture_screens.js)
SCR = os.path.join(HERE, '..', 'docs', 'Attachment_E_Screens')
for k, name in ((3, 'E01_plant_overview.png'), (4, 'E02_11kV_switchboard.png'), (5, 'E07_energy_management.png')):
    shape = d.inline_shapes[k]
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    path = os.path.join(SCR, name)
    d.part.related_parts[rid]._blob = open(path, 'rb').read()
    w, h = Image.open(path).size
    shape.height = int(shape.width * h / w)

# ---------------------------------------------------------------- sweep
BAD = re.compile(r'SICAM|CP-80|Siemens|SITOP|Moxa|MGate|Rittal|Pfannenberg|IEC 60870|IEC 104|SCC\b|\bI/O rows?\b|\brows? \d|CI-85|DI-81|DO-82|AI-83|PS-86|IPC647|OptiPlex')
left = [p.text for p in all_pars() if BAD.search(p.text) and not p.text.startswith('Revision 2 changes the control platform')]
assert not left, left
d.save(OUT)
print('saved', OUT)
