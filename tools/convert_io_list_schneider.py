"""Rev 2 of the Villa Somalia SCADA Device and I/O List: control platform changed to Schneider Electric.

Input is Rev 1 (Siemens SICAM A8000, from tools/harmonize_io_list.py). Output keeps every signal
and tag but allocates the hardwired signals to Modicon M580 X80 and Edge I/O NTS modules at five
I/O locations (tools/schneider_io_allocation.py), and replaces the controller, network, SCADA,
power and enclosure items with Schneider Electric equipment where Schneider makes it.

Usage: python3 tools/convert_io_list_schneider.py <rev1.xlsx> <rev2.xlsx>
"""
import os
import sys
from copy import copy
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from schneider_io_allocation import allocate, MODULES, MODDESC, NAI4  # noqa: E402

SRC, OUT = sys.argv[1], sys.argv[2]
REV = 'Rev 2'
wb = openpyxl.load_workbook(SRC)
LOCNAME = {'SWR': '11 kV switchroom', 'PRS': 'President office station', 'PMS': 'Prime Minister office station',
           'YRD': 'BESS and inverter transformer yard', 'GEN': 'Generator area'}
ROW2LOC = {1: 'SWR', 2: 'SWR', 3: 'PRS', 4: 'PMS', 5: 'YRD', 6: 'GEN'}
TXT = [  # applied to free text in every sheet, in order
    ('I/O rows 3 to 6', 'the remote I/O islands'), ('Rows 3-6', 'PRS, PMS, YRD, GEN'), ('Rows 3, 4', 'PRS, PMS'), ('Rows 5, 6, 2', 'YRD, GEN, SWR'),
    ('Rows 5, 6', 'YRD, GEN'), ('Rows 1, 3, 4', 'SWR, PRS, PMS'), ('Rows 2, 6', 'SWR, GEN'), ('Rows 1, 2', 'SWR'), ('Rows 1 to 6', 'All locations'),
    ('Rows 1-6', 'All locations'), ('Row 1', 'SWR'), ('Row 2', 'SWR'), ('Row 5', 'YRD'), ('Row 6', 'GEN'),
    ('to row 1', 'to the switchroom I/O'), ('to row 2', 'to the switchroom I/O'), ('to row 5', 'to the yard I/O island'), ('to row 6', 'to the generator area I/O island'),
    ('(row 1)', '(switchroom I/O)'), ('(row 2)', '(switchroom I/O)'), ('(row 5)', '(yard I/O island)'), ('(row 6)', '(generator area I/O island)'),
    ('(rows 3, 4)', '(office station I/O islands)'), ('(rows 5, 6)', '(yard and generator area I/O islands)'),
    ('inverter to row 5, DG step-up to row 6, auxiliary to row 2', 'inverter to the yard island, DG step-up to the generator area island, auxiliary to the switchroom'),
    ('nearest I/O row', 'nearest I/O location'), ('wired to row 2', 'wired to the switchroom I/O'),
    ('Not connected to the CP-8050 directly.', 'Not connected to the controllers directly.'),
    ('Only the CP-8050 writes settings to it.', 'Only the controllers write settings to it.'),
    ('Modbus RTU, converted by MGate gateway', 'Modbus RTU, through the Edge I/O serial master (NTSFMB0120K)'),
    ('Modbus TCP, or Modbus RTU through a gateway', 'Modbus TCP, or Modbus RTU through an Edge I/O serial master'),
    ('Modbus RTU through MGate', 'Modbus RTU through the Edge I/O serial master'), ('SICAM CP-8050', 'Modicon M580 controllers'),
]


def fix(s):
    if not isinstance(s, str):
        return s
    for a, b in TXT:
        s = s.replace(a, b)
    return s


def styles_of(ws, row):
    return [copy(ws.cell(row=row, column=c)._style) for c in range(1, ws.max_column + 1)]


def clear(ws, first=1):
    for row in ws.iter_rows(min_row=first, max_row=ws.max_row):
        for c in row:
            c.value = None
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))


def put(ws, r, vals, st):
    for c, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=c, value=v)
        if st is not None and c <= len(st):
            x._style = copy(st[c - 1])


# ---------------------------------------------------------------- IO List: Schneider module and channel
io = wb['IO List']
hdr = [c.value for c in io[1]]
col = {h: i for i, h in enumerate(hdr)}
sigs = [r for r in io.iter_rows(min_row=2, values_only=True) if r[0]]
alloc, N_RELAYS = allocate(sigs, col)
where = {a[0][0]: a for a in alloc}
io.cell(row=1, column=col['I/O row'] + 1, value='I/O location')
io.cell(row=1, column=col['Module type'] + 1, value='Module reference')
last_io = 1
for r in range(2, io.max_row + 1):
    tag = io.cell(row=r, column=1).value
    if not tag:
        continue
    s, mod, ch, iface, relay = where[tag]
    loc, ref, _ = MODULES[mod]
    io.cell(row=r, column=col['I/O row'] + 1, value=loc)
    io.cell(row=r, column=col['Module'] + 1, value=mod)
    io.cell(row=r, column=col['Module type'] + 1, value=ref if ref != NAI4 else 'NTS analogue input (to select)')
    io.cell(row=r, column=col['Channel'] + 1, value=ch)
    for k in ('Description', 'Equipment'):
        x = io.cell(row=r, column=col[k] + 1)
        x.value = fix(x.value)
    n = io.cell(row=r, column=col['Notes'] + 1)
    extra = 'Via interposing relay %s.' % relay if relay else ''
    base = fix(n.value)
    if base and extra and not base.rstrip().endswith('.'):
        base = base.rstrip() + '.'
    n.value = ' '.join(x for x in (base, extra) if x) or None
    last_io = r

# ---------------------------------------------------------------- IO Summary
su = wb['IO Summary']
h_st, b_st = styles_of(su, 1), styles_of(su, 2)
clear(su)
put(su, 1, ['I/O location', 'Location', 'Module', 'Module reference', 'Channels', 'Used', 'Spare', 'Spare %'], h_st)
r = 2
for mid, (loc, ref, chs) in MODULES.items():
    put(su, r, [loc, LOCNAME[loc], mid, ref if ref != NAI4 else 'NTS AI (to select)', chs,
                "=COUNTIF('IO List'!$G$2:$G$%d,C%d)" % (last_io, r), '=E%d-F%d' % (r, r), '=IF(E%d=0,0,G%d/E%d)' % (r, r, r)], b_st)
    su.cell(row=r, column=8).number_format = '0%'
    r += 1
mod_last = r - 1
r += 1
put(su, r, ['Totals by module reference'], None)
su.cell(row=r, column=1).font = copy(su.cell(row=1, column=1).font)
su.cell(row=r, column=1)._style = copy(h_st[0])
r += 1
put(su, r, ['Module reference', 'Description', 'Modules', 'Channels', 'Used', 'Spare', 'Spare %'], h_st)
r += 1
t0 = r
for ref in dict.fromkeys(v[1] for v in MODULES.values()):
    key = ref if ref != NAI4 else 'NTS AI (to select)'
    rng = '$D$2:$D$%d' % mod_last
    put(su, r, [key, MODDESC[ref], '=COUNTIF(%s,A%d)' % (rng, r), '=SUMIF(%s,A%d,$E$2:$E$%d)' % (rng, r, mod_last),
                '=SUMIF(%s,A%d,$F$2:$F$%d)' % (rng, r, mod_last), '=D%d-E%d' % (r, r), '=IF(D%d=0,0,F%d/D%d)' % (r, r, r)], b_st)
    su.cell(row=r, column=7).number_format = '0%'
    r += 1
put(su, r, ['Total', None, '=SUM(C%d:C%d)' % (t0, r - 1), '=SUM(D%d:D%d)' % (t0, r - 1), '=SUM(E%d:E%d)' % (t0, r - 1), '=SUM(F%d:F%d)' % (t0, r - 1),
            '=IF(D%d=0,0,F%d/D%d)' % (r, r, r)], b_st)
su.cell(row=r, column=7).number_format = '0%'
r += 2
put(su, r, ['Hardwired signals by type'], [h_st[0]])
r += 1
put(su, r, ['Type', 'Meaning', 'Signals'], h_st)
r += 1
s0 = r
for typ, meaning in (('DI', 'Digital input (status, alarm)'), ('DO', 'Digital output (command, through an interposing relay)'), ('AI', 'Analogue input (4-20 mA)')):
    put(su, r, [typ, meaning, "=COUNTIF('IO List'!$E$2:$E$%d,A%d)" % (last_io, r)], b_st)
    r += 1
put(su, r, ['Total', None, '=SUM(C%d:C%d)' % (s0, r - 1)], b_st)

# ---------------------------------------------------------------- IO Rows -> IO Locations
lo = wb['IO Rows']
lo.title = 'IO Locations'
h_st, b_st = styles_of(lo, 1), styles_of(lo, 2)
clear(lo)
put(lo, 1, ['I/O location', 'Location', 'I/O hardware', 'Interface reference', 'Link to the controllers', 'Equipment wired here', 'I/O modules', 'Hardwired signals', 'Spare channels'], h_st)
LOCS = [
    ('SWR', '11 kV switchroom, SCADA cabinet', 'Modicon X80 remote drop', 'BMECRA31210 on BMEXBP1200', 'M580 Ethernet remote I/O, copper, in the cabinet',
     '11 kV panels AH1 to AH8, alarm beacon, controller watchdog, auxiliary transformer and LV board, UPS, DC supplies, switchroom E-stop and doors'),
    ('PRS', 'President office transformer station', 'Modicon Edge I/O NTS island', 'NTSNEC1200K', 'EtherNet/IP over the fibre ring, through the station switch', 'RMU, 630 kVA transformer, LVDB, RMU UPS; ST260E relay on the island serial master'),
    ('PMS', 'Prime Minister office transformer station', 'Modicon Edge I/O NTS island', 'NTSNEC1200K', 'EtherNet/IP over the fibre ring, through the station switch', 'RMU, 630 kVA transformer, LVDB, RMU UPS; ST260E relay on the island serial master'),
    ('YRD', 'BESS and inverter transformer yard', 'Modicon Edge I/O NTS island', 'NTSNEC1200K', 'EtherNet/IP over the fibre ring, through field switch 2', 'Inverter transformer, PV and BESS LV breakers, inverter AC breakers, battery fire panel, BESS E-stop and door, panel supervision'),
    ('GEN', 'Generator area', 'Modicon Edge I/O NTS island', 'NTSNEC1200K', 'EtherNet/IP over the fibre ring, through field switch 3', 'DG1 and DG2 status and fuel level, DG step-up transformer and its LV breaker, panel supervision'),
]
for i, v in enumerate(LOCS, 2):
    put(lo, i, list(v) + ["=COUNTIF('IO Summary'!$A$2:$A$%d,A%d)" % (mod_last, i), "=COUNTIF('IO List'!$F$2:$F$%d,A%d)" % (last_io, i),
                          "=SUMIF('IO Summary'!$A$2:$A$%d,A%d,'IO Summary'!$G$2:$G$%d)" % (mod_last, i, mod_last)], b_st)
tr = 2 + len(LOCS)
put(lo, tr, ['Total', None, None, None, None, None, '=SUM(G2:G%d)' % (tr - 1), '=SUM(H2:H%d)' % (tr - 1), '=SUM(I2:I%d)' % (tr - 1)], b_st)
notes = [
    'Both controllers of the M580 hot standby pair reach every location; only the active one reads inputs and drives outputs. No I/O sits in the CPU racks.',
    'Field contacts are wired a short distance to the local I/O. Inputs are dry contacts wetted at 24 V DC from the panel supply; commands drive interposing relays in the 110 V DC control circuit.',
    'Islands at the yard and generator area are in outdoor enclosures (IP55 minimum), shaded or air-conditioned where exposed to direct sun.',
    'Further remote drops or Edge I/O islands can be added for new locations without changing the controllers.',
]
for i, t in enumerate(notes):
    lo.cell(row=tr + 2 + i, column=1, value=t)

# ---------------------------------------------------------------- Device List
dl = wb['Device List']
h = [c.value for c in dl[1]]
dl.cell(row=1, column=h.index('Protocol to CP-8050') + 1, value='Protocol to the controllers')
dl.cell(row=1, column=h.index('Hardwired to I/O row') + 1, value='Hardwired to I/O location')
sec_st = styles_of(dl, 2)
item_st = styles_of(dl, 3)
first_b = next(r for r in range(2, dl.max_row + 1) if str(dl.cell(row=r, column=1).value).startswith('B.'))
for r in range(2, first_b):
    for c in range(1, dl.max_column + 1):
        x = dl.cell(row=r, column=c)
        x.value = fix(x.value)
for r in range(2, first_b):
    if dl.cell(row=r, column=1).value == 'D06':
        dl.cell(row=r, column=11, value='Modbus TCP (IEC 61850 needs an IEC 61850 module on the controllers)')
        dl.cell(row=r, column=12, value='Modbus TCP')
    if dl.cell(row=r, column=1).value == 'D08':
        dl.cell(row=r, column=14, value='Confirm Modbus map with the RMU supplier. Read by the serial master on the office station Edge I/O island.')
old_last = dl.max_row
SE, OFF, TOS, TCS = 'Schneider Electric', 'From the Schneider offer', 'To select', 'To confirm with supplier'
DEV = [
    ('B. Modicon M580 controllers and I/O (Schneider Electric)',),
    ('D18', 'Controller', 'Hot standby CPU, level 2, with remote and distributed I/O', SE, 'Modicon M580', 'BMEH582040', TCS, 2, '11 kV switchroom, SCADA cabinet', 'Ethernet', 'Gateway for all protocols', 'Modbus TCP', '-',
     'Hot standby pair: controller A active, B standby, each in its own rack, linked by the hot standby link. The offer quotes the simplex BMEP582040; requote as hot standby. Schneider to confirm the switchover time.'),
    ('D19', 'Controller', 'Ethernet backplane, 6 slots, dual power supply (CPU racks)', SE, 'Modicon X80', 'BMEXBP0602', TCS, 2, '11 kV switchroom, SCADA cabinet', '-', '-', 'N/A', '-', 'One per CPU. Dual power supply racks are named in the offer for hot standby.'),
    ('D20', 'Controller', 'Redundant power supply, 24-48 V DC (CPU racks)', SE, 'Modicon X80', TOS, TCS, 4, '11 kV switchroom, SCADA cabinet', '-', '-', 'N/A', '-', 'Two per CPU rack. Reference to suit the supply voltage.'),
    ('D21', 'Controller', 'Ethernet communication module, 3 ports (device network)', SE, 'Modicon M580', 'BMENOC0311', OFF, 2, '11 kV switchroom, SCADA cabinet', '3 x RJ45', 'Modbus TCP, EtherNet/IP', 'N/A', '-', 'One per CPU rack: relays, meters, EMS, PLC box, weather logger, Edge I/O islands.'),
    ('D22', 'Controller', 'Ethernet communication module with IP forwarding (control network)', SE, 'Modicon M580', 'BMENOC0321', OFF, 2, '11 kV switchroom, SCADA cabinet', '3 x RJ45', 'Modbus TCP', 'N/A', '-', 'One per CPU rack: link to the SCADA stations; keeps the control and device networks apart.'),
    ('D23', 'Controller', 'Hot standby link transceivers and cable', SE, 'Modicon M580', TOS, TCS, 1, '11 kV switchroom, SCADA cabinet', 'Fibre or copper', 'Hot standby link', 'N/A', '-', 'Between CPU A and CPU B.'),
    ('D24', 'I/O', 'Ethernet backplane, 12 slots (switchroom drop)', SE, 'Modicon X80', 'BMEXBP1200', OFF, 1, '11 kV switchroom, SCADA cabinet', '-', '-', 'N/A', 'SWR', None),
    ('D25', 'I/O', 'Power supply, 24-48 V DC, 32 W (switchroom drop)', SE, 'Modicon X80', 'BMXCPS3020', OFF, 1, '11 kV switchroom, SCADA cabinet', '-', '-', 'N/A', 'SWR', None),
    ('D26', 'I/O', 'Ethernet remote drop adapter', SE, 'Modicon X80', 'BMECRA31210', OFF, 1, '11 kV switchroom, SCADA cabinet', '2 x RJ45', 'M580 Ethernet remote I/O', 'M580 Ethernet remote I/O', 'SWR', 'Switchroom drop on the remote I/O network, reached by both controllers.'),
    ('D27', 'I/O', 'Digital input, 32 channels, 24 V DC', SE, 'Modicon X80', 'BMXDDI3202K', OFF, 4, '11 kV switchroom, SCADA cabinet', 'Pre-wired cable', 'M580 Ethernet remote I/O', 'M580 Ethernet remote I/O', 'SWR', '94 switchroom inputs. Dry contacts wetted at 24 V DC.'),
    ('D28', 'I/O', 'Digital output, 32 channels, 24 V DC', SE, 'Modicon X80', 'BMXDDO3202K', OFF, 1, '11 kV switchroom, SCADA cabinet', 'Pre-wired cable', 'M580 Ethernet remote I/O', 'M580 Ethernet remote I/O', 'SWR', '14 breaker commands, alarm beacon and watchdog, each through an interposing relay.'),
    ('D29', 'I/O', 'Analogue input, 4 isolated channels', SE, 'Modicon X80', 'BMXAMI0410', TCS, 1, '11 kV switchroom, SCADA cabinet', 'Terminal block', 'M580 Ethernet remote I/O', 'M580 Ethernet remote I/O', 'SWR', '110 V and 24 V DC bus voltages, 4-20 mA.'),
    ('D30', 'I/O', 'Pre-wired cable for 32-channel modules, 3 m', SE, 'Modicon X80', 'BMXFCW303', OFF, 5, '11 kV switchroom, SCADA cabinet', '-', '-', 'N/A', 'SWR', None),
    ('D31', 'I/O', 'Serial link module, Modbus RTU', SE, 'Modicon X80', 'BMXNOM0200', OFF, 1, '11 kV switchroom, SCADA cabinet', 'RS485', 'Modbus RTU', 'Modbus RTU', '-', 'For switchroom devices with RS485 only. Confirm use.'),
    ('D32', 'I/O', 'Edge I/O network interface, 2 x RJ45', SE, 'Modicon Edge I/O NTS', 'NTSNEC1200K', OFF, 4, 'Office stations, BESS yard, generator area', '2 x RJ45', 'EtherNet/IP', 'EtherNet/IP', 'PRS, PMS, YRD, GEN', 'One per island.'),
    ('D33', 'I/O', 'Edge I/O power supply module, 24 V DC', SE, 'Modicon Edge I/O NTS', 'NTSPFB1002HK', OFF, 4, 'Office stations, BESS yard, generator area', '-', '-', 'N/A', 'PRS, PMS, YRD, GEN', None),
    ('D34', 'I/O', 'Edge I/O digital input, 16 channels, 24 V DC', SE, 'Modicon Edge I/O NTS', 'NTSDDI1602XK', OFF, 10, 'Office stations, BESS yard, generator area', 'Terminal block', 'EtherNet/IP', 'EtherNet/IP', 'PRS, PMS, YRD, GEN', 'President 3, Prime Minister 3, yard 2, generators 2.'),
    ('D35', 'I/O', 'Edge I/O digital output, 16 channels, 24 V DC, 0.5 A', SE, 'Modicon Edge I/O NTS', 'NTSDDO1602K', OFF, 2, 'President and PM stations', 'Terminal block', 'EtherNet/IP', 'EtherNet/IP', 'PRS, PMS', '8 commands at each office station, through interposing relays.'),
    ('D36', 'I/O', 'Edge I/O analogue input, 4 x 4-20 mA', SE, 'Modicon Edge I/O NTS', TOS, TCS, 1, 'Generator area', 'Terminal block', 'EtherNet/IP', 'EtherNet/IP', 'GEN', 'DG1 and DG2 fuel levels.'),
    ('D37', 'I/O', 'Edge I/O serial master, Modbus RTU', SE, 'Modicon Edge I/O NTS', 'NTSFMB0120K', OFF, 2, 'President and PM stations', 'RS485', 'Modbus RTU', 'Modbus RTU', '-', 'Reads the ST260E RMU relay at each office station. Add one at the generator area if the generator controllers are RS485 only.'),
    ('D38', 'I/O', 'Edge I/O terminal blocks', SE, 'Modicon Edge I/O NTS', 'NTSXTB02030H, NTSXTB18001XH, NTSXTB18001H, NTSXTB12001H', OFF, 22, 'Office stations, BESS yard, generator area', '-', '-', 'N/A', '-', '8 power supply, 10 input, 2 output and 2 serial master blocks.'),
    ('D39', 'I/O', 'Interposing relay with base, 24 V DC coil, contact rated for the 110 V DC control circuit', SE, 'Zelio Relay', TOS, TCS, N_RELAYS, 'Switchroom and office stations', 'Terminal', 'Hardwired', 'Hardwired', 'SWR, PRS, PMS',
     '%d command relays and 2 for the beacon and watchdog. Relay numbers IR01 to IR%02d are on the IO List.' % (N_RELAYS - 2, N_RELAYS)),
    ('C. Network, time and security',),
    ('D40', 'Network', 'Core switch, managed, with SFP fibre ports', SE, 'Modicon managed switch', TOS, TCS, 2, '11 kV switchroom', 'RJ45 + SFP', 'Ethernet, VLANs, ring', 'N/A', '-',
     'Core switch A and B. The offer quotes MCSESM083F23F1, which has copper ports only.'),
    ('D41', 'Network', 'Field and station switch, managed, with SFP fibre ports', SE, 'Modicon managed switch', TOS, TCS, 5, 'MV station, BESS, generators, 2 office stations', 'RJ45 + SFP', 'Ethernet, ring', 'N/A', '-', 'Replaces the unmanaged switch on the existing sheet.'),
    ('D42', 'Network', 'Fibre transceiver, single-mode', SE, 'SFP for the switches above', TOS, TCS, 14, 'All switches', 'LC fibre', '-', 'N/A', '-', '7-switch ring, 2 per link.'),
    ('D43', 'Network', 'Firewall with VPN', 'Fortinet', 'FortiGate Rugged 60F', 'FGR-60F', 'From memory - confirm', 1, '11 kV switchroom', 'RJ45', 'VPN, off by default', 'N/A', '-', 'Not in the Schneider range offered; second unit optional.'),
    ('D44', 'Network', 'GPS time server with antenna', 'Meinberg', 'LANTIME M300', 'M300/GPS', 'From memory - confirm', 1, '11 kV switchroom', 'RJ45', 'NTP', 'NTP', '-', 'Not in the Schneider range offered.'),
    ('D45', 'Network', 'Armoured single-mode fibre, 12 core, with patch panels', 'To suit', '-', 'Length per route', 'TBC', 3, 'Switchroom-President, switchroom-PM, President-PM', 'LC', '-', 'N/A', '-',
     'Third route closes the ring. The yard and generator area islands connect through field switches 2 and 3.'),
    ('D. SCADA stations, power and enclosures',),
    ('D46', 'SCADA', 'SCADA station, industrial panel PC', SE, 'Industrial panel PC', 'HMIP67GD0F88BE7N00', OFF + ' - model to confirm', 2, 'Control room', 'RJ45', 'Modbus TCP', 'Modbus TCP', '-',
     'Station A active, station B hot standby; each is also an operator station.'),
    ('D47', 'SCADA', 'Engineering laptop', 'Dell', 'Latitude 5430 Rugged', 'Configured to order', TCS, 1, 'Control room', 'RJ45', '-', 'N/A', '-', 'Runs EcoStruxure Control Expert and the SCADA build-time licence.'),
    ('D48', 'SCADA', 'Power meter', SE, 'PowerLogic PM5560', 'METSEPM5560', 'From memory - confirm', 3, 'LVDB President, LVDB PM, auxiliary board', 'RJ45', 'Modbus TCP', 'Modbus TCP', '-', 'The offer lists 5 x PM5320; the design keeps 3 x PM5560.'),
    ('D51', 'SCADA', '24 V DC power supply', SE, 'Phaseo', TOS, TCS, 6, 'Switchroom, office stations, BESS yard, generator area', '-', '-', 'N/A', '-',
     'Contact wetting, Edge I/O and X80 power supplies, switches. Input to suit the 110 V DC station supply where available.'),
    ('D52', 'SCADA', 'UPS for the SCADA cabinet, 3000 VA', 'APC (Schneider Electric)', 'Smart-UPS SRT 3000', 'SRT3000RMXLI', 'From memory - confirm', 1, 'Control room', 'RJ45', 'SNMP', 'N/A', '-', 'About 2 hours runtime with battery pack; size at design.'),
    ('D53', 'SCADA', 'SCADA cabinet, 800 x 2000 x 800 mm', SE, 'Spacial SF', 'Size to suit', TCS, 1, '11 kV switchroom', '-', '-', 'N/A', 'SWR', 'Door contact wired to the switchroom I/O.'),
    ('D54', 'SCADA', 'Outdoor I/O panel, IP55 minimum', SE, 'Thalassa PLM', 'Size to suit', TCS, 4, 'Office stations, BESS yard, generator area', '-', '-', 'N/A', 'PRS, PMS, YRD, GEN', 'Shaded or air-conditioned where exposed to direct sun.'),
    ('D55', 'SCADA', 'Alarm sounder and beacon', SE, 'Harmony XVS', 'Model to select', TCS, 1, 'Control room', 'Terminal', 'Hardwired', 'Hardwired', 'SWR', None),
    ('E. Software',),
    ('D56', 'Software', 'SCADA runtime, 4K tags', SE, 'EcoStruxure Machine SCADA Expert', 'HMIEMSE3PRT4KA', OFF, 2, 'SCADA stations', '-', 'Modbus TCP', 'N/A', '-',
     'One per station; about 707 tags. Operator and Admin roles. Redundancy, historian and retention to be confirmed by Schneider.'),
    ('D57', 'Software', 'SCADA build-time licence, 4K tags', SE, 'EcoStruxure Machine SCADA Expert', 'HMIEMSEBT4KA', OFF, 1, 'Engineering laptop', '-', '-', 'N/A', '-', None),
    ('D58', 'Software', 'Controller engineering software', SE, 'EcoStruxure Control Expert', 'By quotation', TCS, 1, 'Engineering laptop', '-', '-', 'N/A', '-', 'Includes the hot standby configuration.'),
    ('D59', 'Software', 'VPN client', 'Fortinet', 'FortiClient', '-', '-', 1, 'Remote users', '-', '-', 'N/A', '-', None),
]
for r in range(first_b, old_last + 1):
    for c in range(1, dl.max_column + 1):
        dl.cell(row=r, column=c).value = None
r = first_b
for v in DEV:
    put(dl, r, list(v) + [None] * (14 - len(v)), sec_st if len(v) == 1 else item_st)
    r += 1
dl_last = r - 1
dl.auto_filter.ref = 'A1:N%d' % dl_last

# ---------------------------------------------------------------- Protocol Summary
ps = wb['Protocol Summary']
h_st, b_st = styles_of(ps, 1), styles_of(ps, 2)
clear(ps)
put(ps, 1, ['Primary protocol', 'Where used', 'Device lines', 'Device quantity'], h_st)
PROT = [
    ('Modbus TCP', 'Controllers to SCADA stations; 11 kV relays, PLC box, EMS, battery via EMS, generator controllers, power meters, weather logger, UPS'),
    ('EtherNet/IP', 'Controllers to the Edge I/O islands at the office stations, BESS yard and generator area'),
    ('M580 Ethernet remote I/O', 'Controllers to the switchroom remote drop and its X80 modules'),
    ('Modbus RTU', 'ST260E RMU relays through the Edge I/O serial masters; RS485 devices in the switchroom'),
    ('Hardwired', 'Status contacts, alarms and commands wired to the I/O; commands through interposing relays'),
    ('Power line communication', 'PV inverters to the PLC box over the AC cables'),
    ('NTP', 'Time synchronisation from the GPS time server'),
    ('Through logger', "Weather sensors to the blue'Log"),
]
for i, (p, w) in enumerate(PROT, 2):
    put(ps, i, [p, w, "=COUNTIF('Device List'!$L$2:$L$%d,A%d)" % (dl_last, i), "=SUMIF('Device List'!$L$2:$L$%d,A%d,'Device List'!$H$2:$H$%d)" % (dl_last, i, dl_last)], b_st)
ps.cell(row=len(PROT) + 3, column=1, value='Supervision protocols: SNMP v3 and syslog from switches, firewall and time server to the SCADA stations. The controllers and the SCADA exchange data by Modbus TCP (EtherNet/IP is an alternative).')

# ---------------------------------------------------------------- Comms Points and Comms Point List: controller pair and I/O locations
cp = wb['Comms Points']
rows_cp = [[cp.cell(row=r, column=c).value for c in range(1, 12)] for r in range(1, cp.max_row + 1)]
st_cp = [styles_of(cp, r) for r in range(1, cp.max_row + 1)]
out_rows, out_st = [], []
for v, st in zip(rows_cp, st_cp):
    if v[0] == 'SICAM CP-8050 internal':
        v[0] = 'Modicon M580 controller internal'
        if isinstance(v[2], str) and v[2].startswith('I/O row '):
            n = int(v[2].split()[2])
            if n == 2:
                continue  # rows 1 and 2 are one switchroom drop
            loc = ROW2LOC[n]
            v[2] = 'I/O link healthy: %s' % {'SWR': 'switchroom remote drop', 'PRS': 'President station island', 'PMS': 'Prime Minister station island',
                                              'YRD': 'BESS yard island', 'GEN': 'generator area island'}[loc]
            v[10] = REV + '. One per I/O location'
        if isinstance(v[10], str) and 'CI-8532' in v[10]:
            v[10] = REV + '. One per I/O location'
    out_rows.append(v)
    out_st.append(st)
last_item = max(i for i, v in enumerate(out_rows) if v[0] and v[2]) + 1
for i, v in enumerate(out_rows[1:last_item], 2):
    v[8] = '=G%d*H%d' % (i, i)
for v in out_rows[last_item:]:
    if v[7] == 'Total':
        v[8] = '=SUM(I2:I%d)' % last_item
    if v[7] == 'Read':
        v[8] = '=SUMIF(D2:D{0},"Read",I2:I{0})+SUMIF(D2:D{0},"Read/Write",I2:I{0})'.format(last_item)
    if v[7] == 'Write':
        v[8] = '=SUMIF(D2:D{0},"Write",I2:I{0})'.format(last_item)
for v in out_rows:
    v[1] = fix(v[1])
for row in cp.iter_rows():
    for c in row:
        c.value = None
for i, (v, st) in enumerate(zip(out_rows, out_st), 1):
    put(cp, i, v, st)

cl = wb['Comms Point List']
LOCTAG = {1: 'SWR', 3: 'PRS', 4: 'PMS', 5: 'YRD', 6: 'GEN'}
rows_cl = [[cl.cell(row=r, column=c).value for c in range(1, 12)] for r in range(1, cl.max_row + 1)]
st_cl = [styles_of(cl, r) for r in range(1, cl.max_row + 1)]
new_cl, new_st = [], []
for v, st in zip(rows_cl, st_cl):
    if isinstance(v[1], str) and v[1].startswith('SICAM CP-8050'):
        v[1] = v[1].replace('SICAM CP-8050 internal', 'Modicon M580 controller').replace('SICAM CP-8050', 'Modicon M580 controller')
        if isinstance(v[0], str) and v[0].startswith('CP-ROW'):
            n = int(v[0][6])
            if n == 2:
                continue
            loc = LOCTAG[n]
            v[0] = 'CP-IO-%s-OK' % loc
            v[2] = 'I/O link healthy: %s' % {'SWR': 'switchroom remote drop', 'PRS': 'President station island', 'PMS': 'Prime Minister station island',
                                              'YRD': 'BESS yard island', 'GEN': 'generator area island'}[loc]
            v[10] = REV + '. One per I/O location'
    new_cl.append(v)
    new_st.append(st)
for v in new_cl:
    v[6] = fix(v[6])
for row in cl.iter_rows():
    for c in row:
        c.value = None
for i, (v, st) in enumerate(zip(new_cl, new_st), 1):
    put(cl, i, v, st)
n_comms = sum(1 for v in new_cl[1:] if v[0])
cl.auto_filter.ref = 'A1:K%d' % (n_comms + 1)
tags = [v[0] for v in new_cl[1:] if v[0]]
assert len(tags) == len(set(tags))

# ---------------------------------------------------------------- Datasheets
ds = wb['Datasheets']
h_st, b_st = styles_of(ds, 1), styles_of(ds, 2)
keep = [[ds.cell(row=r, column=c).value for c in range(1, 10)] for r in range(2, ds.max_row + 1) if ds.cell(row=r, column=2).value == 'Field device']
clear(ds)
put(ds, 1, ['No.', 'Group', 'Item', 'Manufacturer', 'Model', 'Order no. to search', 'Documents to collect', 'Where to download', 'Collected'], h_st)
SEDL = "se.com: search the reference, tab 'Documents' (datasheet, user manual)"
DSH = [
    ('Controller', 'Hot standby CPU', SE, 'Modicon M580', 'BMEH582040', 'Datasheet; M580 hot standby system planning guide; cybersecurity guide'),
    ('Controller', 'Dual power supply backplane', SE, 'Modicon X80', 'BMEXBP0602', 'Datasheet'),
    ('Controller', 'Redundant power supply', SE, 'Modicon X80', 'To select', 'Datasheet'),
    ('Controller', 'Ethernet communication modules', SE, 'Modicon M580', 'BMENOC0311, BMENOC0321', 'Datasheet; communication modules user manual'),
    ('I/O', 'Ethernet backplane and remote drop adapter', SE, 'Modicon X80', 'BMEXBP1200, BMECRA31210', 'Datasheet; remote I/O system planning guide'),
    ('I/O', 'Power supply', SE, 'Modicon X80', 'BMXCPS3020', 'Datasheet'),
    ('I/O', 'Digital input and output modules', SE, 'Modicon X80', 'BMXDDI3202K, BMXDDO3202K, BMXFCW303', 'Datasheet; discrete I/O user manual'),
    ('I/O', 'Analogue input module', SE, 'Modicon X80', 'BMXAMI0410', 'Datasheet; analogue I/O user manual'),
    ('I/O', 'Serial link module', SE, 'Modicon X80', 'BMXNOM0200', 'Datasheet'),
    ('I/O', 'Edge I/O network interface and power supply', SE, 'Modicon Edge I/O NTS', 'NTSNEC1200K, NTSPFB1002HK', 'Datasheet; Edge I/O NTS user manual'),
    ('I/O', 'Edge I/O digital input and output', SE, 'Modicon Edge I/O NTS', 'NTSDDI1602XK, NTSDDO1602K', 'Datasheet'),
    ('I/O', 'Edge I/O analogue input', SE, 'Modicon Edge I/O NTS', 'To select', 'Datasheet'),
    ('I/O', 'Edge I/O serial master', SE, 'Modicon Edge I/O NTS', 'NTSFMB0120K', 'Datasheet'),
    ('I/O', 'Interposing relays', SE, 'Zelio Relay', 'To select', 'Datasheet incl. DC contact rating'),
    ('SCADA', 'SCADA stations', SE, 'Industrial panel PC', 'HMIP67GD0F88BE7N00', 'Datasheet'),
    ('Software', 'SCADA software', SE, 'EcoStruxure Machine SCADA Expert', 'HMIEMSE3PRT4KA, HMIEMSEBT4KA', 'Product description; system requirements; redundancy and licensing'),
    ('Software', 'Controller engineering software', SE, 'EcoStruxure Control Expert', 'By quotation', 'Product information; system requirements'),
    ('Network', 'Core, field and station switches with SFP ports', SE, 'Modicon managed switch', 'To select', 'Datasheet'),
    ('Network', 'Fibre transceivers', SE, 'SFP, single-mode', 'To select', 'Datasheet'),
    ('Network', 'Firewall and VPN', 'Fortinet', 'FortiGate Rugged 60F', 'FGR-60F', 'Datasheet'),
    ('Network', 'GPS time server', 'Meinberg', 'LANTIME M300/GPS', 'M300/GPS', 'Datasheet'),
    ('Power', 'Power meters', SE, 'PowerLogic PM5560', 'METSEPM5560', 'Datasheet; Modbus register list'),
    ('Power', '24 V DC power supplies', SE, 'Phaseo', 'To select', 'Datasheet'),
    ('Power', 'UPS', 'APC (Schneider Electric)', 'Smart-UPS SRT 3000 VA', 'SRT3000RMXLI', 'Datasheet'),
    ('Enclosure', 'SCADA cabinet', SE, 'Spacial SF', 'Size to suit', 'Datasheet'),
    ('Enclosure', 'Outdoor I/O panels', SE, 'Thalassa PLM', 'Size to suit', 'Datasheet'),
    ('Enclosure', 'Alarm sounder and beacon', SE, 'Harmony XVS', 'Model to select', 'Datasheet'),
]
WHERE = {'Fortinet': 'fortinet.com data sheets library', 'Meinberg': 'meinbergglobal.com, product page LANTIME M300', 'APC (Schneider Electric)': "apc.com or se.com, search 'SRT3000RMXLI'"}
r = 2
for g, it, mf, model, ref, docs in DSH:
    put(ds, r, [r - 1, g, it, mf, model, ref, docs, WHERE.get(mf, SEDL), None], b_st)
    r += 1
for v in keep:
    v[0] = r - 1
    put(ds, r, v, b_st)
    r += 1
ds.cell(row=r + 1, column=1, value="Fill the Collected column as each file is saved. Name files '<No.>_<Model>.pdf' so they sort in this order inside the zip.")

# ---------------------------------------------------------------- Read Me
rm = wb['Read Me']
SET = {
    'Status': 'Rev 2, 16 Oct 2026: control platform changed to Schneider Electric (Modicon M580 hot standby, Modicon X80 and Edge I/O NTS, EcoStruxure Machine SCADA Expert). Signals and tags as Rev 1. For client submission as Attachment B.',
    'Controller': 'Two Schneider Electric Modicon M580 hot standby CPUs (A active, B standby) as plant controller and gateway. The embedded EMS keeps dispatch of PV, battery and generators.',
    'SCADA': 'EcoStruxure Machine SCADA Expert on two SCADA stations (active and hot standby), linked to the controllers by Modbus TCP.',
    'Device List': 'Every device in the design: model, product or ordering number, quantity, location, interface and protocol to the controllers.',
    'IO List': 'Every hardwired signal: tag, description, type, I/O location, Schneider module and channel, alarm level and who may operate it.',
    'IO Locations': 'The five I/O locations: the switchroom remote drop and four Edge I/O islands, their link to the controllers and the equipment wired to each.',
    'IO Summary': 'Channels used and spare per I/O module, calculated from the IO List.',
    'Checked': 'Product number found on a supplier listing or configuration document.',
}
for row in rm.iter_rows():
    for c in row:
        if c.value == 'IO Rows':
            c.value = 'IO Locations'
for row in rm.iter_rows():
    for c in row:
        if c.column == 1 and c.value in SET:
            rm.cell(row=c.row, column=2).value = SET[c.value]
ASSUME = {
    '1': 'Field contacts are dry contacts, wetted at 24 V DC from a field supply in each I/O panel. Breaker and switch commands drive interposing relays whose contacts switch the 110 V DC control circuit; the 24 V DC outputs never switch a coil directly.',
    '3': 'The embedded EMS offers Modbus TCP. The PV inverter data reaches the controllers from the power line communication box, as shown on the existing communication sheet.',
    '6': 'Spare capacity target is about 20% per I/O location.',
    '7': 'Field signals are collected at the nearest I/O location: switchroom remote drop (SWR), President and Prime Minister office stations (PRS, PMS), BESS and inverter transformer yard (YRD), generator area (GEN). The four islands connect over the fibre ring. Actual cable distances are not known yet.',
}
ah = next(c.row for row in rm.iter_rows() for c in row if c.value == 'Assumptions')
r = ah + 1
while rm.cell(row=r, column=1).value is not None:
    k = str(rm.cell(row=r, column=1).value)
    if k in ASSUME:
        rm.cell(row=r, column=2).value = ASSUME[k]
    r += 1
# status key for references taken from the offer
key_row = next(c.row for row in rm.iter_rows() for c in row if c.value == 'TBC' and c.column == 1)
rm.insert_rows(key_row + 1)
rm.cell(row=key_row + 1, column=1, value=OFF)
rm.cell(row=key_row + 1, column=2, value="Reference taken from Schneider Electric's Offer Proposal Villa Somalia, version 1, 16 Oct 2026. Confirm in the requote.")
for c in (1, 2):
    rm.cell(row=key_row + 1, column=c)._style = copy(rm.cell(row=key_row, column=c)._style)
# Revision 1 notes stay for the record, with their Siemens part names; Revision 2 notes follow
nr = rm.max_row + 2
rm.cell(row=nr, column=1, value='Revision 2 changes')
rev1_head = next(c for row in rm.iter_rows() for c in row if c.value == 'Revision 1 changes')
rev1_head.value = 'Revision 1 changes (Siemens platform; replaced by Revision 2 where they differ)'
rm.cell(row=nr, column=1)._style = copy(rev1_head._style)
R2 = [
    'Control platform changed from Siemens SICAM A8000 to Schneider Electric: Modicon M580 hot standby CPUs, X80 remote drop in the switchroom, Edge I/O NTS islands at the office stations, BESS yard and generator area, EcoStruxure Machine SCADA Expert on two SCADA stations.',
    'Same 244 hardwired signals and tags. I/O rows 1 to 6 become five I/O locations (rows 1 and 2 are one switchroom drop); module and channel columns now give the Schneider allocation, %d interposing relays for the commands.' % N_RELAYS,
    'Communication points: the six I/O row link points become five I/O location link points (CP-IO-SWR-OK to CP-IO-GEN-OK); controller points keep their tags. Total %d communication points.' % n_comms,
    'Controller to SCADA link is Modbus TCP (was IEC 60870-5-104). ST260E relays are read by the Edge I/O serial masters instead of MGate gateways. 11 kV relays on Modbus TCP; IEC 61850 needs an added controller module.',
    'Device List sections B to E rebuilt for Schneider equipment (new IDs D18 to D59). Firewall, GPS time server and engineering laptop stay with other makers. IO Rows sheet renamed IO Locations; IO Summary, Protocol Summary and Datasheets rebuilt.',
]
for i, t in enumerate(R2):
    rm.cell(row=nr + 1 + i, column=1, value=i + 1)
    rm.cell(row=nr + 1 + i, column=2, value=t)
    for c in (1, 2):
        rm.cell(row=nr + 1 + i, column=c)._style = copy(rm.cell(row=rev1_head.row + 1, column=c)._style)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('hardwired', len(alloc), 'comms', n_comms, 'relays', N_RELAYS, 'device rows', dl_last, '->', OUT)
